import json, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree
S=json.load(open('stats.json')); G=json.load(open('prisma_groups.json')); P=S['prisma']
NAVY=RGBColor(0x0A,0x25,0x4E); WHITE=RGBColor(255,255,255); INK=RGBColor(0x11,0x18,0x27); MUTED=RGBColor(0x4B,0x55,0x63)
ACC=RGBColor(0xC2,0x41,0x0C); GRAY=RGBColor(0x9C,0xA3,0xAF); LIGHT=RGBColor(0xF3,0xF4,0xF6); RED=RGBColor(0xB9,0x1C,0x1C); GREEN=RGBColor(0x04,0x78,0x57)
FONT='Times New Roman'
prs=Presentation('/root/work/aq.pptx'); s=prs.slides[0]
keep={'Picture 76'}
for sh in list(s.shapes):
    if sh.name not in keep: sh._element.getparent().remove(sh._element)
def tb(x,y,w,h,fill=None,line=None,anchor=MSO_ANCHOR.TOP,margin=0.12,shape=MSO_SHAPE.RECTANGLE):
    if fill is None and line is None and shape==MSO_SHAPE.RECTANGLE:
        b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    else:
        b=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
        b.shadow.inherit=False
        if fill is None: b.fill.background()
        else: b.fill.solid(); b.fill.fore_color.rgb=fill
        if line is None: b.line.fill.background()
        else: b.line.color.rgb=line; b.line.width=Pt(2.5)
        if shape==MSO_SHAPE.ROUNDED_RECTANGLE: b.adjustments[0]=0.08
    tf=b.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for a in ('margin_left','margin_right'): setattr(tf,a,Inches(margin))
    tf.margin_top=tf.margin_bottom=Inches(0.06)
    return b
def write(b,paras,size=20,color=INK,align=PP_ALIGN.LEFT,space=4,line_spacing=None):
    """paras: list of items; item = str | list of (text,bold,italic) | dict(runs=..., bullet=bool, size=..)"""
    tf=b.text_frame; first=True
    for it in paras:
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        if isinstance(it,dict): runs=it['runs']; bullet=it.get('bullet',False); sz=it.get('size',size); al=it.get('align',align); col=it.get('color',color)
        else: runs=it; bullet=False; sz=size; al=align; col=color
        if isinstance(runs,str): runs=[(runs,False,False)]
        p.alignment=al; p.space_after=Pt(space)
        if line_spacing: p.line_spacing=line_spacing
        if bullet:
            pPr=p._p.get_or_add_pPr(); pPr.set('marL',str(int(Inches(0.32)))); pPr.set('indent',str(-int(Inches(0.28))))
            bu=etree.SubElement(pPr,qn('a:buFont')); bu.set('typeface','Arial')
            bc=etree.SubElement(pPr,qn('a:buChar')); bc.set('char','•')
        for t,bo,itl in runs:
            r=p.add_run(); r.text=t; f=r.font; f.name=FONT; f.size=Pt(sz); f.bold=bo; f.italic=itl; f.color.rgb=col
    return b
def header(x,y,w,text,h=0.95,size=44):
    b=tb(x,y,w,h,fill=NAVY,anchor=MSO_ANCHOR.MIDDLE); write(b,[text],size=size,color=WHITE,align=PP_ALIGN.CENTER); return b
def body(x,y,w,h,paras,size=20,**k):
    b=tb(x,y,w,h,fill=WHITE,margin=0.18); write(b,paras,size=size,**k); b.text_frame.margin_top=Inches(0.12); return b
# ---------------- TITLE ----------------
band=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,Inches(6.85)); band.fill.solid(); band.fill.fore_color.rgb=WHITE; band.line.fill.background(); band.shadow.inherit=False
sp=s.shapes._spTree; sp.remove(band._element); sp.insert(2,band._element)
t=tb(6.6,0.25,40.8,6.4,anchor=MSO_ANCHOR.MIDDLE)
write(t,[{'runs':[('Psychosocial Outcomes and Coping Mechanisms Among Women with',True,False)],'size':66,'align':PP_ALIGN.CENTER,'color':NAVY},
         {'runs':[('Premenstrual Dysphoric Disorder: A Global Scoping Review',True,False)],'size':66,'align':PP_ALIGN.CENTER,'color':NAVY},
         {'runs':[('Ifeoluwanimi P. Shobayo, BSc, MSPHc',False,False)],'size':36,'align':PP_ALIGN.CENTER}],space=6)
# ---------------- LEFT COLUMN ----------------
LX,LW=0.5,11.3
header(LX,7.05,LW,'Abstract')
body(LX,8.0,LW,7.25,[
 [('Background: ',True,False),('Premenstrual dysphoric disorder (PMDD) is a DSM-5 depressive disorder affecting an estimated 1.6% of women and girls (~31 million) worldwide, yet its psychosocial burden—especially in low- and middle-income countries (LMICs)—has not been mapped.',False,False)],
 [('Methods: ',True,False),("Scoping review (Arksey & O'Malley; PRISMA-ScR) of seven databases, 2010–2025: peer-reviewed empirical studies with PMDD as the population and a psychosocial outcome or coping mechanism as a primary aim.",False,False)],
 [('Results: ',True,False),(f"{P['identified_total']:,} records yielded {P['included']} studies. Depression (n = 28), mood/emotion regulation (n = 18), psychiatric comorbidity (n = 18), distress (n = 16), suicidality (n = 13) and functioning (n = 13) dominated; coping was studied in 7. Two-thirds of studies came from high-income countries; 3 came from sub-Saharan Africa.",False,False)],
 [('Conclusions: ',True,False),('Evidence is growing but geographically concentrated, relies on provisional diagnoses and rarely examines coping or lived experience.',False,False)],
 [('Keywords: ',True,False),('PMDD; psychosocial outcomes; coping; suicidality; LMICs; sub-Saharan Africa',False,True)]],size=24,space=7)
header(LX,15.5,LW,'Introduction and Research Question')
body(LX,16.45,LW,8.2,[
 'PMDD is a severe, hormone-linked mood disorder in which irritability, mood swings, depressed mood and anxiety recur in the late luteal phase and remit after menses begin. It is classified as a depressive disorder in DSM-5 and recognised in ICD-11.',
 'Using prospectively confirmed diagnoses, community prevalence is about 1.6%—roughly 31 million women and girls (Reilly et al., 2024)—and premenstrual symptoms are common across Africa (pooled PMS prevalence 47%; Andualem et al., 2024).',
 'Because symptoms recur every month, PMDD can affect mental health, relationships, school and work. This evidence is scattered across disciplines, and it is unclear where in the world it has been generated.',
 [('Research Question',True,False)],
 [('What psychosocial outcomes and coping mechanisms have been documented among women with PMDD, and how is this evidence distributed across global settings—particularly LMICs and sub-Saharan Africa?',False,True)]],size=25,space=12)
header(LX,24.9,LW,'Methods')
rows=[('Study design',"Scoping review (Arksey & O'Malley, 2005; Levac et al., 2010), reported per PRISMA-ScR (Tricco et al., 2018)."),
('Databases','CINAHL Ultimate, APA PsycInfo, APA PsycArticles, MEDLINE Ultimate, Women’s Studies International, PubMed, Cochrane Library.'),
('Years & language','2010–2025; English.'),
('Inclusion','Peer-reviewed empirical studies; PMDD (provisional or confirmed) as the defined population; ≥1 psychosocial outcome or coping mechanism as a primary aim.'),
('Exclusion','Animal, laboratory/biological, intervention, review, instrument-validation, case-report and non-peer-reviewed formats; PMS-only samples.'),
('Screening & charting','Deduplication (DOI + title); title/abstract screening with logged reasons; charted design, country, World Bank income group (FY2027), PMDD ascertainment and 13 psychosocial domains.')]
gt=s.shapes.add_table(len(rows)+1,2,Inches(LX),Inches(25.85),Inches(LW),Inches(9.6)).table
gt.columns[0].width=Inches(3.0); gt.columns[1].width=Inches(LW-3.0)
for j,h in enumerate(['Category','Description']):
    c=gt.cell(0,j); c.fill.solid(); c.fill.fore_color.rgb=RGBColor(0xD9,0xE2,0xF3)
    c.text_frame.paragraphs[0].text=''; r=c.text_frame.paragraphs[0].add_run(); r.text=h; r.font.bold=True; r.font.size=Pt(20); r.font.name=FONT; r.font.color.rgb=INK
for i,(a,b_) in enumerate(rows,1):
    for j,txt in enumerate((a,b_)):
        c=gt.cell(i,j); c.fill.solid(); c.fill.fore_color.rgb=WHITE; c.vertical_anchor=MSO_ANCHOR.MIDDLE
        p=c.text_frame.paragraphs[0]; p.text=''; r=p.add_run(); r.text=txt; r.font.size=Pt(20); r.font.name=FONT; r.font.bold=(j==0); r.font.color.rgb=INK
        c.margin_left=c.margin_right=Inches(0.1); c.margin_top=c.margin_bottom=Inches(0.05)
# remove table style banding look
gt.rows[0].height=Inches(0.6)
for i in range(1,len(rows)+1): gt.rows[i].height=Inches((9.6-0.6)/len(rows))
tblPr=gt._tbl.tblPr; tblPr.set('bandRow','0'); tblPr.set('firstRow','1')
# ---------------- CENTER PANEL ----------------
CX,CW=12.2,24.6
panel=tb(CX,7.05,CW,28.45,fill=WHITE)
# --- Fig 1 PRISMA (native shapes) ---
fx,fy=12.6,7.35
write(tb(fx,fy,11.2,0.6),[[('Fig 1: PRISMA-ScR Flow of Records',True,False)]],size=22,color=NAVY)
def pbox(x,y,w,h,lines,fill=WHITE,line=NAVY,size=16,bold_first=False,align=PP_ALIGN.CENTER):
    b=tb(x,y,w,h,fill=fill,line=line,anchor=MSO_ANCHOR.MIDDLE,shape=MSO_SHAPE.ROUNDED_RECTANGLE,margin=0.12)
    items=[]
    for k,l in enumerate(lines): items.append({'runs':[(l,bold_first and k==0,False)],'align':align})
    write(b,items,size=size,space=1); return b
def arrow(x1,y1,x2,y2):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=NAVY; c.line.width=Pt(3)
    ln=c.line._get_or_add_ln(); te=etree.SubElement(ln,qn('a:tailEnd')); te.set('type','triangle'); te.set('w','med'); te.set('len','med')
def stage(y,h,label):
    b=tb(fx,y,0.75,h,fill=NAVY,anchor=MSO_ANCHOR.MIDDLE,shape=MSO_SHAPE.ROUNDED_RECTANGLE,margin=0.02)
    b.text_frame.word_wrap=False
    bp=b.text_frame._txBody.find(qn('a:bodyPr')); bp.set('vert','vert270')
    write(b,[{'runs':[(label,True,False)],'align':PP_ALIGN.CENTER}],size=16,color=WHITE)
stage(8.2,2.3,'Identification'); stage(10.9,6.3,'Screening'); stage(17.6,2.3,'Included')
bx,bw=13.6,4.9; rx,rw=18.95,4.85
pbox(bx,8.2,bw,2.3,['Records identified from databases',f"(n = {P['identified_total']:,})"],size=17,bold_first=True)
db=P['identified_by_database']
pbox(rx,8.2,rw,2.3,[f"{k}: {v:,}" for k,v in db.items()],fill=LIGHT,line=GRAY,size=14,align=PP_ALIGN.LEFT)
arrow(bx+bw/2,10.5,bx+bw/2,10.9)
pbox(bx,10.9,bw,1.6,['Records after duplicates removed',f"(n = {P['screened']:,})"],size=16)
pbox(rx,10.9,rw,1.6,[f"Duplicates removed (n = {P['duplicates_removed']})"],fill=LIGHT,line=GRAY,size=15)
arrow(bx+bw,11.7,rx,11.7)
arrow(bx+bw/2,12.5,bx+bw/2,13.3)
pbox(bx,13.3,bw,2.2,['Records screened','(title, abstract, subject terms)',f"(n = {P['screened']:,})"],size=16)
short={'PMDD not the primary study population':'PMDD not primary population','Published outside 2010–2025':'Outside 2010–2025','Reviews, meta-analyses, guidelines, commentaries':'Reviews/commentaries','Other ineligible publication types (books, chapters, dissertations, letters, editorials, case reports, conference abstracts)':'Other ineligible formats','Non-English full text':'Non-English','Intervention or treatment study':'Intervention studies','Laboratory/biological study':'Laboratory/biological','No psychosocial outcome or coping mechanism as a primary aim':'No psychosocial/coping aim','Instrument development/validation study':'Instrument validation','Animal study':'Animal studies'}
ex=sorted(G.items(),key=lambda x:-x[1])
pbox(rx,12.9,rw,4.3,[f"Records excluded (n = {P['excluded_total']:,})"]+[f"{short[k]}: {v}" for k,v in ex],fill=RGBColor(0xFE,0xF2,0xF2),line=RED,size=15,bold_first=True,align=PP_ALIGN.LEFT)
arrow(bx+bw,14.4,rx,14.4)
arrow(bx+bw/2,15.5,bx+bw/2,17.6)
pbox(bx,17.6,bw,2.3,['Studies included',f"(n = {P['included']})"],fill=RGBColor(0xEC,0xFD,0xF5),line=GREEN,size=20,bold_first=True)
write(tb(rx,17.5,rw,2.5),[{'runs':[('Eligibility judged on titles, abstracts and subject terms; full-text review ongoing. Source: authors’ screening of database exports.',False,True)]}],size=12,color=MUTED)
# --- Fig 2 domains (native bar chart) ---
dx=24.2
write(tb(dx,fy,12.3,0.6),[[('Fig 2: Psychosocial Domains Examined (n = 85 studies)',True,False)]],size=22,color=NAVY)
D=sorted(S['domains'].items(),key=lambda x:x[1])
lab={'Depression / depressive symptoms':'Depression','Psychological distress & perceived stress':'Distress & perceived stress','Suicidal ideation, attempts & self-harm':'Suicidality & self-harm','Interpersonal & relational functioning':'Interpersonal functioning','Academic & occupational functioning':'Academic/occupational functioning','Lived experience, stigma & help-seeking':'Lived experience & help-seeking','Trauma, adversity & discrimination':'Trauma & adversity'}
cd=CategoryChartData(); cd.categories=[lab.get(k,k) for k,_ in D]; cd.add_series('Studies',[v for _,v in D])
gf=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(dx),Inches(8.0),Inches(12.3),Inches(12.0),cd); ch=gf.chart
def style_chart(ch,size=15,maxv=None):
    ch.has_title=False; ch.has_legend=False; ch.font.name='Arial'; ch.font.size=Pt(size); ch.font.color.rgb=INK
    pl=ch.plots[0]; pl.gap_width=45; pl.has_data_labels=True; dl=pl.data_labels; dl.font.size=Pt(size); dl.font.name='Arial'; dl.position=XL_LABEL_POSITION.OUTSIDE_END; dl.number_format='0'; dl.number_format_is_linked=False
    ser=pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb=NAVY
    va=ch.value_axis; va.has_major_gridlines=False; va.visible=False
    if maxv: va.maximum_scale=maxv
    va.minimum_scale=0
    ca=ch.category_axis; ca.tick_labels.font.size=Pt(size); ca.format.line.color.rgb=GRAY; ca.has_major_gridlines=False
    return ser
style_chart(ch,15,32)
write(tb(dx,20.0,12.3,0.8),[{'runs':[('Studies may address more than one domain. Source: authors’ charting of included studies.',False,True)]}],size=12,color=MUTED)
# --- stat tiles ---
ty=21.1
tiles=[(str(P['included']),'studies included\n(2010–2025)'),('67%','from high-income\ncountries'),('3','studies from\nsub-Saharan Africa'),('24%','confirmed PMDD with\nprospective ratings')]
tw=(CW-0.8-3*0.35)/4
for k,(big,small) in enumerate(tiles):
    x=CX+0.4+k*(tw+0.35)
    b=tb(x,ty,tw,3.1,fill=RGBColor(0xEE,0xF2,0xF8),anchor=MSO_ANCHOR.MIDDLE,shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    col=ACC if k==2 else NAVY
    write(b,[{'runs':[(big,True,False)],'size':60,'align':PP_ALIGN.CENTER,'color':col},{'runs':[(small,False,False)],'size':19,'align':PP_ALIGN.CENTER,'color':INK}],space=0)
# --- Fig 3 income + region ---
gy=24.55
write(tb(12.6,gy,11.2,0.6),[[('Fig 3: Studies by World Bank Income Group (FY2027)',True,False)]],size=22,color=NAVY)
order=['High income','Upper-middle income','Lower-middle income','Low income','Not classifiable (online/multinational/not reported)']
names=['High','Upper-middle','Lower-middle','Low','Online / NR']
cd=CategoryChartData(); cd.categories=names; cd.add_series('Studies',[S['income'].get(k,0) for k in order])
gf=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(12.6),Inches(gy+0.7),Inches(11.2),Inches(9.2),cd); ch=gf.chart
ser=style_chart(ch,16,64)
for i,colr in enumerate([NAVY,ACC,ACC,ACC,GRAY]):
    pt=ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb=colr
write(tb(12.6,34.5,11.2,0.9),[{'runs':[('Orange = LMICs (21 studies, 25%); only 5 from low/lower-middle-income countries.',False,True)]}],size=13,color=MUTED)
write(tb(dx,gy,12.3,0.6),[[('Fig 4: Studies by World Bank Region',True,False)]],size=22,color=NAVY)
R=sorted(S['region'].items(),key=lambda x:x[1])
rl={'Online / multinational / not reported':'Online / multinational / NR','Europe & Central Asia':'Europe & Central Asia'}
cd=CategoryChartData(); cd.categories=[rl.get(k,k) for k,_ in R]; cd.add_series('Studies',[v for _,v in R])
gf=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(dx),Inches(gy+0.7),Inches(12.3),Inches(9.2),cd); ch=gf.chart
ser=style_chart(ch,16,33)
for i,(k,_) in enumerate(R):
    if k=='Sub-Saharan Africa': pt=ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb=ACC
write(tb(dx,34.5,12.3,0.9),[{'runs':[('Sub-Saharan Africa: Ethiopia (n = 2), Nigeria (n = 1). Europe & Central Asia includes Türkiye.',False,True)]}],size=13,color=MUTED)
# ---------------- RIGHT COLUMN ----------------
RX,RW=37.1,10.5
header(RX,7.05,RW,'Results')
B=lambda t: {'runs':t if isinstance(t,list) else [(t,False,False)],'bullet':True}
body(RX,8.0,RW,10.45,[
 B([('Growing evidence: ',True,False),('49% of studies published since 2020; 56% cross-sectional, only 3 qualitative.',False,False)]),
 B([('Diagnosis: ',True,False),('only 24% confirmed PMDD with prospective daily ratings; 42% used provisional screening.',False,False)]),
 B([('Suicidality: ',True,False),('PMDD linked with suicidal ideation (OR 2.22) and attempts (OR 2.10) in a US national sample; 72% lifetime active ideation and 34% attempts in a global PMDD sample.',False,False)]),
 B([('Depression: ',True,False),('2.6-fold higher risk of later depression in a Taiwanese national cohort.',False,False)]),
 B([('Quality of life & relationships: ',True,False),('lower QoL (Sweden, n = 17,284) and more relationship disruption (IRR 1.22).',False,False)]),
 B([('Trauma & adversity: ',True,False),('adverse childhood experiences were linked to PMDD in Iceland and Japan (aOR 5.61 for ≥4 ACEs).',False,False)]),
 B([('Coping: ',True,False),('only 7 studies—mostly self-medication (analgesics, cannabis) or maladaptive styles.',False,False)]),
 B([('Geography: ',True,False),('67% high-income; 5 studies (6%) from low/lower-middle-income countries; 3 (3.5%) from sub-Saharan Africa; 1 LMIC study used prospective diagnosis.',False,False)])],size=26,space=10)
header(RX,18.7,RW,'Discussion and Conclusion',size=40)
body(RX,19.65,RW,5.55,[
 'Evidence on PMDD’s psychosocial burden is expanding and consistent—depression, suicidality, impaired functioning and reduced quality of life—but it is concentrated in high-income countries, relies on provisional diagnoses and rarely examines coping or lived experience.',
 'In sub-Saharan Africa the burden is largely undocumented: no study examined suicidality, quality of life or coping. Culturally responsive screening, integration of PMDD into sexual, reproductive and mental health services, and Africa-centred research investment are needed.'],size=27,space=12)
header(RX,25.4,RW,'Limitations',size=40)
body(RX,26.35,RW,2.5,[B('Eligibility judged on titles/abstracts; full-text review ongoing.'),B('English only; African Journals Online and grey literature not searched.'),B('Many studies combined PMS and PMDD or used provisional diagnoses.')],size=24,space=5)
header(RX,29.05,RW,'Future Work',size=40)
body(RX,30.0,RW,2.3,[B('Prospectively confirmed PMDD studies in LMICs and sub-Saharan Africa.'),B('Qualitative research on stigma, lived experience and coping.'),B('Validate PMDD tools in African languages.')],size=24,space=5)
header(RX,32.5,RW,'References',size=34,h=0.7)
refs=["1. Reilly TJ et al. J Affect Disord. 2024;349:534-540.",
"2. Andualem F et al. Front Psychiatry. 2024;15:1338304.",
"3. Arksey H, O'Malley L. Int J Soc Res Methodol. 2005;8(1):19-32.",
"4. Tricco AC et al. Ann Intern Med. 2018;169(7):467-473.",
"5. Pilver CE et al. Soc Psychiatry Psychiatr Epidemiol. 2013;48(3):437-446.",
"6. Eisenlohr-Moul T et al. BMC Psychiatry. 2022;22. doi:10.1186/s12888-022-03851-0",
"7. Li DJ et al. Asian J Psychiatr. 2023;79:103355.",
"8. Wang Q et al. JAMA Netw Open. 2025;8(9):e2533823.",
"9. Westermark V et al. J Affect Disord. 2024;364:132-138.",
"Full list of 85 included studies: see manuscript and supplementary file."]
body(RX,33.2,RW,2.3,[{'runs':[(r,False,False)]} for r in refs],size=12,space=0)
prs.save('PMDD_APHA_Poster.pptx'); print('saved')
