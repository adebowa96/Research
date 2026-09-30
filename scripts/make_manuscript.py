"""Builds the full manuscript (manuscript/MRKH_Scoping_Review_Manuscript.docx).
Run make_figures.py and make_map.js first. Usage: python3 scripts/make_manuscript.py

Inline markup: **bold**, *italic*, ^{sup} superscript, [[...]] = yellow-highlighted
placeholder the authors must complete or confirm before submission."""
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import manuscript_results
from citations import Citer

HERE = Path(__file__).parent
ROOT = HERE.parent
FIG = ROOT / "figures"
data = json.loads((HERE / "data.json").read_text())
N = data["total_included"]
P = data["prisma"]
OUT = dict(data["outcomes"])
DES = {name: (n, note) for name, n, note in data["designs"]}
import csv as _csv
EXTRACTION = list(_csv.DictReader(open(ROOT / "rescreen" / "extraction_provisional.csv", encoding="utf-8")))
REG = {name: n for name, n, _ in data["regions"]}


def pct(n, d=N, digits=1):
    return f"{100 * n / d:.{digits}f}%"


doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
normal.paragraph_format.space_after = Pt(0)
for lvl, size in ((1, 14), (2, 12), (3, 12)):
    st = doc.styles[f"Heading {lvl}"]
    st.font.name = "Times New Roman"
    rfonts = st.element.rPr.rFonts
    rfonts.set(qn("w:eastAsia"), "Times New Roman")
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)  # theme fonts would override Times New Roman
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.italic = lvl == 3
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

# page numbers in footer
fp = doc.sections[0].footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
    run = fp.add_run()
    if tag:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tag)
    else:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = text
    run._r.append(el)

TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|\^\{.+?\}|\[\[.+?\]\])")
BODY_WORDS = []  # words in main text (Introduction through Conclusions) for the word count
counting = {"on": False}


CITER = Citer()


def add_runs(par, text, size=None, bold=False):
    text = CITER.resolve(text).replace("{N_INCLUDED}", str(N))
    for piece in TOKEN.split(text):
        if not piece:
            continue
        b, i, sup, hl = bold, False, False, False
        if piece.startswith("**"):
            piece, b = piece[2:-2], True
        elif piece.startswith("[["):
            piece, hl = "[" + piece[2:-2] + "]", True
        elif piece.startswith("^{"):
            piece, sup = piece[2:-1], True
        elif piece.startswith("*"):
            piece, i = piece[1:-1], True
        r = par.add_run(piece)
        r.bold, r.italic = b, i
        if sup:
            r.font.superscript = True
        if hl:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if size:
            r.font.size = Pt(size)
        if counting["on"] and not sup:
            BODY_WORDS.extend(piece.split())


def para(text, align=None, indent=True, size=None, bold=False, style=None, spacing=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if indent and style is None and align is None:
        p.paragraph_format.first_line_indent = Inches(0.5)
    if spacing is not None:
        p.paragraph_format.line_spacing_rule = spacing
    add_runs(p, text, size=size, bold=bold)
    return p


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        add_runs(p, it)


def heading(text, level=1):
    h = doc.add_heading(level=level)
    if pending_break["on"]:
        h.paragraph_format.page_break_before = True
        pending_break["on"] = False
    add_runs(h, text)
    return h


pending_break = {"on": False}


def page_break():
    # applied as "page break before" on the next heading, so a full page never
    # leaves a stray break paragraph that creates a blank page
    pending_break["on"] = True


def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def table(caption, header, rows, widths=None, note=None, size=10):
    was, counting["on"] = counting["on"], False  # tables are excluded from the word count
    cp = doc.add_paragraph()
    cp.paragraph_format.keep_with_next = True
    cp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    cp.paragraph_format.space_before = Pt(12)
    cp.paragraph_format.space_after = Pt(4)
    add_runs(cp, caption)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.text = ""
        add_runs(c.paragraphs[0], h, size=size, bold=True)
        set_cell_shading(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ""
            add_runs(cells[j].paragraphs[0], str(v), size=size)
    for row in t.rows:
        for j, c in enumerate(row.cells):
            for p in c.paragraphs:
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                p.paragraph_format.first_line_indent = None
            if widths:
                c.width = Inches(widths[j])
    if widths:
        for j, col in enumerate(t.columns):
            col.width = Inches(widths[j])
    if note:
        np_ = doc.add_paragraph()
        np_.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        np_.paragraph_format.space_after = Pt(12)
        add_runs(np_, note, size=10)
    counting["on"] = was
    return t


def figure(path, caption, width=6.0):
    was, counting["on"] = counting["on"], False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph()
    cp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    cp.paragraph_format.space_after = Pt(12)
    add_runs(cp, caption, size=10)
    counting["on"] = was


# ====================================================================== TITLE PAGE
para("**Mental Health and Psychosocial Outcomes Among Individuals With "
     "Mayer-Rokitansky-Küster-Hauser (MRKH) Syndrome: A Scoping Review**",
     align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
para("**Running title:** Mental health in MRKH syndrome: a scoping review",
     align=WD_ALIGN_PARAGRAPH.CENTER)
para("Ifeoluwanimi P. Shobayo, BSc, MSPHc^{1}; Paul Okojie, PhD^{1}; Robyn Anderson, PhD^{1}",
     align=WD_ALIGN_PARAGRAPH.CENTER)
para("^{1}Department of Public and Community Health, Liberty University, Lynchburg, Virginia, USA",
     align=WD_ALIGN_PARAGRAPH.CENTER)
para("", indent=False)
para("**Corresponding author:** Ifeoluwanimi P. Shobayo, Department of Public and Community "
     "Health, Liberty University, Lynchburg, VA, USA. Email: [[corresponding author email]]",
     indent=False)
para("**Word count (main text):** {WORDCOUNT}   **Tables:** 4   **Figures:** 4   "
     "**Supplementary material:** Appendices A–C", indent=False)
para("**Prior presentation:** An abstract of this work was submitted to the American Public "
     "Health Association (APHA) 2026 Annual Meeting & Expo, Sexual and Reproductive Health "
     "Section, and accepted for poster presentation.",
     indent=False)
para("**Keywords:** Mayer-Rokitansky-Küster-Hauser syndrome; Müllerian agenesis; mental health; "
     "coping; psychosocial outcomes; quality of life; reproductive health", indent=False)
page_break()

# ====================================================================== ABSTRACT
heading("Abstract")
abstract = [
    "**Background:** Mayer-Rokitansky-Küster-Hauser (MRKH) syndrome is a rare congenital "
    "condition, affecting approximately 1 in 4,500–5,000 female births, characterized by "
    "uterovaginal agenesis in individuals with a 46,XX karyotype and functional ovaries. "
    "Typically diagnosed during adolescence, MRKH is associated with significant psychosocial "
    "burden, yet mental health outcomes and coping mechanisms remain insufficiently synthesized, "
    "limiting their integration into clinical care and public health planning. This scoping "
    "review mapped the mental health outcomes, coping mechanisms, and healthcare system gaps "
    "reported among individuals with MRKH.",
    "**Methods:** Following Arksey and O'Malley's framework and the PRISMA Extension for Scoping "
    "Reviews (PRISMA-ScR), we searched PubMed, Scopus, and (via EBSCOhost) MEDLINE, CINAHL, APA "
    "PsycInfo, and Women's Studies International for studies published between January 2019 "
    "and March 2026. Eligible studies were primary quantitative, "
    "qualitative, and mixed-methods research reporting psychological outcomes and/or coping "
    "mechanisms in MRKH populations.",
    f"**Results:** Of {P['identified']} records identified ({P['screened']} after "
    f"deduplication), {N} studies were included. Psychosexual and relational "
    f"challenges were the most frequently reported outcomes (n = "
    f"{OUT['Psychosexual & relational challenges']}; "
    f"{pct(OUT['Psychosexual & relational challenges'], digits=0)}), followed by reduced quality "
    f"of life, body image, and self-esteem (n = {OUT['QoL, body image & self-esteem']}; "
    f"{pct(OUT['QoL, body image & self-esteem'], digits=0)}), broader psychological distress "
    f"(n = {OUT['Broader psychological distress']}; "
    f"{pct(OUT['Broader psychological distress'], digits=0)}), and depression and anxiety "
    f"(n = {OUT['Depression & anxiety']}; {pct(OUT['Depression & anxiety'], digits=0)}). Coping "
    f"mechanisms were documented in {OUT['Coping mechanisms documented']} studies, mostly peer "
    f"support and self-management, and healthcare system gaps in {OUT['Healthcare system gaps']}. "
    f"Studies came mainly from Europe (n = {REG['Europe']}) and Asia (n = {REG['Asia']}).",
    "**Conclusions:** MRKH-related psychosocial burden is substantial yet under-integrated into "
    "care models. Multidisciplinary, mental health–inclusive care and greater research "
    "investment, particularly in underrepresented and low-resource settings, are urgently needed.",
]
for a in abstract:
    para(a, indent=False)
page_break()

# ====================================================================== MAIN TEXT
counting["on"] = True
heading("Introduction")
para("Mayer-Rokitansky-Küster-Hauser (MRKH) syndrome, also termed Müllerian agenesis or "
     "Müllerian aplasia, is a congenital condition characterized by absence of the uterus and "
     "upper two-thirds of the vagina in individuals with a 46,XX karyotype, functional ovaries, "
     "and typical secondary sexual development. It affects approximately 1 in 4,500–5,000 "
     "female births^{@acog}; a nationwide Danish registry study estimated a prevalence of 1 in 4,982 "
     "live female births.^{@herlin16} MRKH is classified as type I (isolated uterovaginal aplasia) or "
     "type II (accompanied by extragenital anomalies, most often renal, skeletal, auditory, or "
     "cardiac).^{@herlin20} Diagnosis is usually made in adolescence during evaluation for primary "
     "amenorrhea,^{@acog,@herlin20} a developmental period that is critical for identity formation, body "
     "image, and social relationships.")
para("The diagnosis carries profound implications for fertility, psychosexual development, and "
     "self-image. Management commonly includes vaginal dilation as first-line treatment, surgical "
     "creation of a neovagina when indicated, and fertility options such as gestational "
     "surrogacy and, more recently, uterus transplantation.^{@acog,@herlin20} The American College of "
     "Obstetricians and Gynecologists identifies psychosocial counseling as one of the most "
     "important components of effective management.^{@acog} Earlier studies documented elevated "
     "psychological distress and lower self-esteem among women with MRKH compared with "
     "controls,^{@hb09} higher anxiety scores among affected adolescents,^{@laggari} and poorer "
     "mental health–related quality of life, higher anxiety, and impaired sexual wellness "
     "relative to normative data.^{@liao} A review of this early literature concluded that creating a "
     "neovagina alone does not ensure a good psychological outcome and that psychological "
     "support at critical times can be helpful,^{@bean} while qualitative work described how young "
     "women manage intimacy, a sensitivity to difference, and threats to identity after "
     "diagnosis.^{@patterson} More recent qualitative studies describe emotional turmoil at diagnosis, "
     "challenges to sexual identity and intimate relationships, the profound impact of "
     "infertility,^{@R011} and reduced self-esteem and sexual wellbeing.^{@R017}")
para("Two systematic reviews have since synthesized the psychological, quality-of-life, and "
     "sexual outcomes of MRKH.^{@facchin,@tsarna} However, to our knowledge, the more recent literature has "
     "not been mapped with a public health lens that jointly considers mental health outcomes, "
     "the coping mechanisms individuals use, and the healthcare system gaps that shape access to "
     "psychosocial support. As a result, psychosocial outcomes remain insufficiently integrated "
     "into clinical and public health frameworks for this population.")
para("A scoping review is well suited to mapping a heterogeneous body of quantitative, "
     "qualitative, and mixed-methods evidence and identifying gaps for future research and "
     "practice.^{@arksey,@tricco} This scoping review therefore addressed the following research question: "
     "*What mental health outcomes and coping mechanisms are reported among individuals with MRKH "
     "syndrome, and what healthcare system gaps exist?*")

heading("Methods")
heading("Study Design", 2)
para("This scoping review followed Arksey and O'Malley's five-stage methodological "
     "framework^{@arksey}—(1) identifying the research question; (2) identifying relevant studies; "
     "(3) selecting studies; (4) charting the data; and (5) collating, summarizing, and reporting "
     "the results—and is reported in accordance with the PRISMA Extension for Scoping Reviews "
     "(PRISMA-ScR; Appendix C),^{@tricco} informed by published guidance on the steps of conducting a "
     "scoping review.^{@mak} A review protocol was not registered.")
heading("Eligibility Criteria", 2)
para("Eligibility was defined using the Population–Concept–Context approach. The *population* "
     "was individuals diagnosed with MRKH syndrome; the *concepts* were psychological or mental "
     "health outcomes and coping mechanisms; and the *context* was any healthcare or community "
     "setting in any country. We included primary quantitative, qualitative, and mixed-methods "
     "studies published between January 2019 and March 2026 that reported psychological "
     "outcomes and/or coping mechanisms among individuals with MRKH. Studies focused solely on "
     "anatomical, surgical, or fertility outcomes were excluded, as were records published "
     "before January 2019 and records not published in English (the English-only criterion "
     "recorded in the review's March 2026 screening tracker). Reviews, commentaries, "
     "editorials, and case reports were excluded as not primary research.")
heading("Information Sources and Search Strategy", 2)
para("Searches were run and exported on March 29, 2026, in PubMed, Scopus, and EBSCOhost "
     "[[confirm the PubMed search was also run on this date]]. "
     "The EBSCOhost search covered MEDLINE Ultimate, CINAHL Ultimate, APA PsycInfo, and Women's "
     "Studies International, and APA PsycInfo was also exported separately (Appendix A). Date "
     "limits were not applied identically across interfaces: the Scopus search was limited to "
     "2019–2026, the PubMed export covered 2021–2026 (MEDLINE records from 2019–2020 were "
     "captured through the EBSCOhost MEDLINE search), and EBSCOhost results were exported "
     "without a date limit. Records published before January 2019 were therefore removed during "
     "title and abstract screening. Search strategies combined controlled "
     "vocabulary and free-text terms for MRKH (e.g., “Mayer-Rokitansky-Küster-Hauser,” “MRKH,” "
     "“Müllerian agenesis,” “vaginal agenesis”) with terms for mental health and coping (e.g., "
     "“depression,” “anxiety,” “psychological distress,” “quality of life,” “body image,” "
     "“coping”). Records retrieved from each source are reported in Appendix A. [[Verify that "
     "the example terms above match the search you ran; recover exact strings from your "
     "PubMed, Scopus, and EBSCOhost search histories if saved.]]")
manuscript_results.build(dict(para=para, heading=heading, table=table, figure=figure, P=P, N=N,
                               OUT=OUT, DES=DES, REG=REG, FIG=FIG, pct=pct, data=data,
                               extraction=EXTRACTION))
counting["on"] = False

# ====================================================================== DECLARATIONS
heading("Declarations")
decl = [
    "**Acknowledgements:** This scoping review was conducted under the mentorship of Dr. Paul "
    "Okojie and Dr. Robyn Anderson, Department of Public and Community Health, Liberty University.",
    "**Funding:** No external funding was received.",
    "**Conflicts of interest:** [[The authors declare no conflicts of interest. — confirm with "
    "all authors.]]",
    "**Author contributions:** [[e.g., I.P.S.: conceptualization, search, screening, data "
    "charting, analysis, writing—original draft; P.O. and R.A.: supervision, methodology, "
    "writing—review and editing. Edit to reflect actual contributions (CRediT taxonomy).]]",
    "**Ethics approval:** Not applicable; this review analyzed published studies.",
    "**Data availability:** The data charting form and extracted data are available from the "
    "corresponding author on reasonable request.",
    "**Use of AI tools:** Generative AI (Claude, Anthropic) was used to propose title and "
    "abstract screening decisions, chart data from abstracts, draft and edit the manuscript, "
    "verify reference details, and prepare figures. The authors "
    "reviewed and verified all content and take full responsibility for the work. [[Confirm "
    "wording against the target journal's AI policy.]]",
]
for dtext in decl:
    para(dtext, indent=False)

# ====================================================================== REFERENCES
heading("References")
refs = CITER.reference_list()
missing = [e["record_id"] for e in EXTRACTION if e["record_id"] not in CITER.order]
assert not missing, f"included studies not cited in text: {missing}"
for i, r in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.35)
    add_runs(p, f"{i}.\t{r}")

# ====================================================================== APPENDICES
page_break()
heading("Appendix A. Search Strategy")
para("Searches were exported on March 29, 2026. Record counts below are taken from the export "
     "files and total 516, matching Figure 1. [[Paste the exact search string for each database "
     "if it can be recovered from your search history.]]", indent=False)
para("[[Suggested PubMed/MEDLINE string for comparison. Replace it with the string you actually "
     "ran; do not report a search you did not run:]]", indent=False)
para("(\"Mayer-Rokitansky-Kuster-Hauser\"[tiab] OR \"Mayer-Rokitansky-Küster-Hauser\"[tiab] OR "
     "MRKH[tiab] OR Rokitansky[tiab] OR \"Mullerian agenesis\"[tiab] OR \"Mullerian aplasia\"[tiab] "
     "OR \"vaginal agenesis\"[tiab] OR \"uterovaginal agenesis\"[tiab] OR "
     "\"Mullerian Ducts/abnormalities\"[Mesh]) AND (depress*[tiab] OR anxiety[tiab] OR "
     "\"mental health\"[tiab] OR psycholog*[tiab] OR distress[tiab] OR \"quality of life\"[tiab] "
     "OR \"body image\"[tiab] OR \"self-esteem\"[tiab] OR coping[tiab] OR psychosocial[tiab] OR "
     "\"Mental Health\"[Mesh] OR \"Adaptation, Psychological\"[Mesh] OR \"Quality of Life\"[Mesh]) "
     "AND (\"2019/01/01\"[dp] : \"2026/03/31\"[dp])", indent=False, size=10)
para("[[Scopus: TITLE-ABS-KEY(...) with the same two concept blocks; PsycINFO and CINAHL: the "
     "same keywords plus each database's own subject headings (e.g., APA Thesaurus “Coping "
     "Behavior”; CINAHL Headings “Quality of Life”).]]", indent=False)
table("**Table A1.** Records retrieved by database",
      ["Database", "Platform", "Date searched", "Records retrieved"],
      [["MEDLINE Ultimate; CINAHL Ultimate; Women's Studies International; APA PsycInfo "
        "(combined export: 145; 39; 15; 8)", "EBSCOhost", "Mar 29, 2026", "207"],
       ["APA PsycInfo (separate export)", "EBSCOhost", "Mar 29, 2026", "20"],
       ["MEDLINE (publications 2021–2026)", "PubMed", "[[Mar 29, 2026?]]", "146"],
       ["Scopus (publications 2019–2026)", "Scopus", "Mar 29, 2026", "143"],
       ["**Total**", "", "", f"**{P['identified']}**"]],
      widths=[2.9, 1.1, 1.2, 1.2],
      note="Records published before 2019 (100 exported records; 89 after deduplication, all "
           "from EBSCOhost) were excluded at title and abstract screening.")

page_break()
heading("Appendix B. Characteristics of Included Studies")
para(f"Data charting was checked against the full text for {data['full_text_checked']} of {N} studies; "
     "entries marked † were charted from the abstract or title only. "
     "Domains: D, depression/anxiety; Q, quality of life, body image, self-esteem; S, "
     "psychosexual and relational; P, broader psychological distress; C, coping documented; "
     "H, healthcare system gaps. Reference numbers refer to the main reference list.",
     indent=False, size=10, spacing=WD_LINE_SPACING.SINGLE)
dom_cols = [("Depression & anxiety", "D"), ("Reduced QoL, body image & self-esteem", "Q"),
            ("Psychosexual & relational challenges", "S"), ("Broader psychological distress", "P"),
            ("Coping mechanisms documented", "C"), ("Healthcare system gaps", "H")]
table("**Table B1.** Characteristics of included studies (N = {N_INCLUDED})",
      ["Study", "Country", "Design", "n", "Domains", "Coping"],
      [[CITER.resolve(f"{e['study']}^{{@{e['record_id']}}}") + ("†" if e["full_text_checked"] == "no" else ""), e["country"],
        e["design_detail"], e["sample_n"],
        " ".join(code for col, code in dom_cols if e[col] == "yes"), e["coping_types"] or "—"]
       for e in sorted(EXTRACTION, key=lambda e: e["study"])],
      widths=[1.1, 0.9, 1.5, 0.8, 0.7, 1.5], size=8)

page_break()
heading("Appendix C. PRISMA-ScR Checklist")
prisma_items = [
    ("Title", "1", "Title page"),
    ("Structured summary", "2", "Abstract"),
    ("Rationale", "3", "Introduction"),
    ("Objectives", "4", "Introduction (research question)"),
    ("Protocol and registration", "5", "Methods: Study Design [[complete]]"),
    ("Eligibility criteria", "6", "Methods: Eligibility Criteria"),
    ("Information sources", "7", "Methods: Information Sources; Appendix A"),
    ("Search", "8", "Appendix A [[complete]]"),
    ("Selection of sources of evidence", "9", "Methods: Study Selection"),
    ("Data charting process", "10", "Methods: Data Charting"),
    ("Data items", "11", "Methods: Data Charting"),
    ("Critical appraisal of individual sources of evidence (optional)", "12",
     "Methods: Synthesis (not performed)"),
    ("Synthesis of results", "13", "Methods: Synthesis"),
    ("Selection of sources of evidence", "14", "Results; Figure 1"),
    ("Characteristics of sources of evidence", "15", "Results; Table 1; Appendix B"),
    ("Critical appraisal within sources of evidence (optional)", "16", "Not applicable"),
    ("Results of individual sources of evidence", "17", "Appendix B"),
    ("Synthesis of results", "18", "Results; Tables 2–4; Figures 2–4"),
    ("Summary of evidence", "19", "Discussion: Summary of Evidence"),
    ("Limitations", "20", "Discussion: Strengths and Limitations"),
    ("Conclusions", "21", "Discussion: Conclusions"),
    ("Funding", "22", "Declarations"),
]
table("**Table C1.** PRISMA-ScR checklist (Tricco et al., 2018)",
      ["Section / topic", "Item", "Reported in"],
      [[a, b, c] for a, b, c in prisma_items], widths=[3.0, 0.6, 2.8])

# fill in the word count on the title page
wc = len(BODY_WORDS)
for p in doc.paragraphs:
    for r in p.runs:
        if "{WORDCOUNT}" in r.text:
            r.text = r.text.replace("{WORDCOUNT}", f"{wc:,}")

out = ROOT / "manuscript" / "MRKH_Scoping_Review_Manuscript.docx"
out.parent.mkdir(exist_ok=True)
doc.save(out)
print(f"wrote {out} (main text ≈ {wc} words)")

# ====================================================================== SEPARATE REFERENCE LIST
ref_doc = Document()
for sec in ref_doc.sections:
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
st = ref_doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
doc_backup, doc = doc, ref_doc  # reuse para()/add_runs() helpers on the reference document
counting["on"] = False
para("**References (AMA 11th edition)**", align=WD_ALIGN_PARAGRAPH.CENTER, size=14,
     spacing=WD_LINE_SPACING.SINGLE)
para("*Mental Health and Psychosocial Outcomes Among Individuals With "
     "Mayer-Rokitansky-Küster-Hauser (MRKH) Syndrome: A Scoping Review*",
     align=WD_ALIGN_PARAGRAPH.CENTER, spacing=WD_LINE_SPACING.SINGLE)
para("Numbered in order of first citation in the manuscript. In-text citations are superscript "
     "numerals, placed after periods and commas and before colons and semicolons.",
     indent=False, size=10, spacing=WD_LINE_SPACING.SINGLE)
for i, r in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.35)
    p.paragraph_format.space_after = Pt(6)
    add_runs(p, f"{i}.\t{r}")
para("**Poster reference list** (poster numbering; shortened AMA format)", indent=False,
     spacing=WD_LINE_SPACING.SINGLE)
poster_map = ["acog", "herlin16", "arksey", "tricco", "hb09", "laggari", "liao", "facchin"]
for i, m in enumerate(poster_map, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.35)
    add_runs(p, f"{i}.\t{CITER.lib[m]}")
doc = doc_backup
ref_out = ROOT / "manuscript" / "References_AMA.docx"
ref_doc.save(ref_out)
print("wrote", ref_out)
