"""Build the AJPH-style manuscript aligned with the approved abstract and the APHA poster (44 studies)."""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------- reference pool (poster numbering 1-51 from the AMA reference document) ----------
REFDOC = '/home/user/Research/PMDD_Scoping_Review/PMDD_Poster_References_AMA.docx'
POOL = {}
for p in Document(REFDOC).paragraphs:
    m = re.match(r'^(\d+)\.\t(.*)$', p.text)
    if m: POOL[int(m.group(1))] = m.group(2).strip()
assert len(POOL) == 51

order = []  # poster numbers in order of first citation
def num(k):
    if k not in order: order.append(k)
    return order.index(k) + 1
def cite_str(keys):
    ns = sorted(num(k) for k in keys)
    out, i = [], 0
    while i < len(ns):
        j = i
        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1: j += 1
        if j - i >= 2: out.append(f'{ns[i]}–{ns[j]}')
        else: out += [str(x) for x in ns[i:j + 1]]
        i = j + 1
    return ','.join(out)

# ---------- document setup ----------
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
pf = st.paragraph_format; pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE; pf.space_after = Pt(0)
for s in doc.sections: s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
fp = doc.sections[0].footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run()
for tag, val in (('begin', None), ('instr', 'PAGE'), ('end', None)):
    if tag == 'instr':
        it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = val; r._r.append(it)
    else:
        f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), tag); r._r.append(f)

NOTE = RGBColor(0xB4, 0x53, 0x09)
def runs(p, text, size=None, bold=False, italic=False, color=None):
    """Markup: {1,2} = citation (poster ref numbers), **bold**, *italic*, [[note]] = author note."""
    for tok in re.split(r'(\{[\d,\s]+\}|\*\*.+?\*\*|\*.+?\*|\[\[.+?\]\])', text):
        if not tok: continue
        b, it, sup, col = bold, italic, False, color
        if tok.startswith('{'):
            tok = cite_str([int(x) for x in tok[1:-1].split(',')]); sup = True
        elif tok.startswith('**'): tok = tok[2:-2]; b = True
        elif tok.startswith('[['): tok = '[' + tok[2:-2] + ']'; col = NOTE; b = True
        elif tok.startswith('*'): tok = tok[1:-1]; it = True
        rr = p.add_run(tok); rr.bold = b; rr.italic = it; rr.font.superscript = sup
        if size: rr.font.size = Pt(size)
        if col: rr.font.color.rgb = col
def P(text, indent=True, align=None, keep=False, size=None, single=False):
    p = doc.add_paragraph()
    if indent: p.paragraph_format.first_line_indent = Inches(0.5)
    if align == 'center': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if single: p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if keep: p.paragraph_format.keep_with_next = True
    runs(p, text, size=size); return p
def H1(t):
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
    rr = p.add_run(t.upper()); rr.bold = True
def H2(t):
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
    rr = p.add_run(t); rr.bold = True; rr.italic = True
def page_break(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
WORDS = []  # main-text paragraphs for the word count
def T(text): WORDS.append(text); return P(text)

# ---------- title page ----------
P('**Psychosocial Outcomes and Coping Mechanisms Among Women With Premenstrual Dysphoric Disorder: A Global Scoping Review**', indent=False, align='center')
P('Ifeoluwanimi P. Shobayo, MSPH, Chelsea R. Mazonde, MPH, Marylyn O. Oduneye, MSPH, Tahirou Diallo, MSPH, Fadzai G. Nyarugwe, MSPH, and Cynthia C. Ilechukwu, MSPH', indent=False, align='center')
P('[[Add each author’s department, institution, city, state and country; the corresponding author’s mailing address, email and telephone; and ORCID iDs.]]', indent=False)
P('**Keywords:** premenstrual dysphoric disorder; psychosocial outcomes; coping mechanisms; global health; scoping review', indent=False)
wc_par = P('', indent=False)  # word counts filled in at the end
page_break()

# ---------- abstract ----------
H1('Abstract')
ABS = [
 ('Objectives', 'To map psychosocial outcomes and coping mechanisms among women with premenstrual dysphoric disorder (PMDD) and the geographic distribution of this evidence.'),
 ('Methods', 'Following Arksey and O’Malley’s framework and PRISMA-ScR, we searched 5 databases for peer-reviewed studies (2010–2025) with PMDD as the primary population and a psychosocial outcome or coping mechanism as a primary aim.'),
 ('Results', 'Of 1,943 records, 1,618 were screened after deduplication, 64 were assessed for eligibility and 44 were included. Depression was most documented (n = 39), followed by psychological distress (n = 24), interpersonal functioning (n = 18), suicidal ideation or self-harm (n = 14) and quality of life (n = 14). Only 3 studies examined coping mechanisms. Three studies (6.8%) came from low- and lower-middle-income countries; Nigeria was the only sub-Saharan African setting.'),
 ('Conclusions', 'Evidence on the psychosocial burden of PMDD is growing but concentrated in high-income countries, and how women cope remains largely unexamined.'),
 ('Public Health Implications', 'Suicide risk screening, integrated mental and reproductive health services and research in African and other low-income settings are needed.'),
]
abs_words = 0
for h, t in ABS:
    P(f'***{h}.*** {t}'.replace('***', '**'), indent=False); abs_words += len(t.split()) + len(h.split())
page_break()

# ---------- introduction ----------
T('Premenstrual dysphoric disorder (PMDD) is a recurring, hormone-linked mood disorder in which irritability, mood swings, depressed mood and anxiety emerge in the late luteal phase of the menstrual cycle and remit shortly after menses begin. It is classified as a depressive disorder in the *Diagnostic and Statistical Manual of Mental Disorders, Fifth Edition* (DSM-5),{2} and it is recognized in the *International Classification of Diseases, 11th Revision*.{3} Its diagnostic validity is still debated, in part because many studies rely on retrospective recall rather than prospective daily symptom ratings.{20}')
T('How common PMDD appears depends on how it is diagnosed. A 2024 meta-analysis estimated a pooled prevalence of 3.2% when diagnosis was confirmed with prospective ratings and 7.7% when it was provisional; in community samples with confirmed diagnosis, prevalence was 1.6%, roughly 31 million women and girls worldwide.{1} Premenstrual symptoms more broadly are common in Africa, where the pooled prevalence of premenstrual syndrome reaches 47%.{8}')
T('Because PMDD recurs every month across the reproductive years, its consequences extend beyond symptom counts. PMDD has been linked with suicidal ideation and attempts,{4,5} disrupted romantic relationships{6} and reduced quality of life.{7} People with PMDD also describe misdiagnosis, medical mistrust and uneven provider knowledge.{18,19} This evidence is spread across psychiatry, gynecology, nursing and psychology journals, and it is unclear which psychosocial outcomes have been studied, how women cope and where the evidence has been generated. That question matters for public health: a condition whose burden has not been documented in a setting is unlikely to be recognized, screened for or funded there.')
T('We therefore conducted a scoping review to map psychosocial outcomes and coping mechanisms reported among women with PMDD and to identify geographic gaps in the published literature. The review asked: What psychosocial outcomes and coping mechanisms have been reported in studies of women with PMDD published between 2010 and 2025, and how are these studies distributed across geographic settings?')

# ---------- methods ----------
H1('Methods')
H2('Study Design')
T('This scoping review followed the framework of Arksey and O’Malley{9} with the refinements proposed by Levac et al{10} and is reported according to the Preferred Reporting Items for Systematic Reviews and Meta-Analyses extension for Scoping Reviews (PRISMA-ScR).{11} Consistent with scoping review methodology, the aim was to map the extent and nature of the evidence rather than to appraise study quality or pool effect estimates. [[State whether a protocol was registered (e.g., Open Science Framework, with DOI and date) or that no protocol was registered.]]')
H2('Eligibility Criteria')
T('Eligible studies were peer-reviewed empirical quantitative or qualitative studies published between 2010 and 2025 that had PMDD as the primary population and examined at least one psychosocial outcome or coping mechanism as a primary aim. Psychosocial outcomes included depression, psychological distress (including anxiety and perceived stress), interpersonal functioning, suicidal ideation or self-harm and quality of life. No restriction was placed on country or setting. We excluded animal studies, laboratory research, reviews, intervention studies and non-peer-reviewed publications.')
H2('Information Sources and Search')
T('We searched PubMed/MEDLINE, CINAHL, APA PsycInfo, APA PsycArticles and the Cochrane Library. Search strategies combined terms for premenstrual dysphoric disorder (e.g., “premenstrual dysphoric disorder,” “PMDD”) with psychosocial and coping concepts, and results were limited to 2010–2025. [[Give the date each database was searched and put the full search string for every database in a supplemental appendix.]]')
H2('Selection of Sources of Evidence')
T('Records from all databases were combined, and duplicates were identified by digital object identifier and normalized title. After removing duplicates, we screened titles and abstracts against the eligibility criteria. Records that passed this step were assessed for eligibility against the full criteria using their titles, abstracts and database subject terms. Full texts were not reviewed at this stage. [[Describe who screened records (e.g., 2 independent reviewers), how disagreements were resolved and, if any software or AI tool assisted screening or writing, disclose it according to AJPH policy.]]')
H2('Data Charting and Synthesis')
T('For each included study, we charted the first author, publication year, country setting and the psychosocial domains examined. Country was taken from the study setting described in the abstract, and countries were grouped using the World Bank FY2027 income classification.{12} Studies that used online or multinational samples, or whose setting could not be determined from the available records, were not assigned to a country. Each study was coded to every psychosocial domain it examined, so domain counts are not mutually exclusive. Publication years were grouped into 4-year periods. We summarized the results as frequencies and a narrative synthesis organized by domain and setting.')

# ---------- results ----------
H1('Results')
H2('Study Selection and Characteristics')
T('The search identified 1,943 records. After 325 duplicates were removed, 1,618 records were screened by title and abstract, and 1,554 were excluded because they were not focused on PMDD, did not have a psychosocial or coping aim, or were reviews, intervention studies, or animal or laboratory studies. Of the 64 records assessed for eligibility, 20 did not meet the inclusion criteria, and 44 studies were included (Figure 1; Table 1).')
TABLE_CALLOUT = len(order)  # table citations are numbered here, at the first callout of Table 1
TABLE_ROWS = [  # (study, poster ref, year, setting, income group)
 ('Roy et al', 15, 2025, 'Bangladesh', 'Lower-middle'), ('Hodgetts and Kinghorn', 21, 2025, 'United Kingdom', 'High'),
 ('Gordon et al', 22, 2025, 'Not assigned', '—'), ('Torabi et al', 23, 2025, 'Not assigned', '—'),
 ('Habib et al', 19, 2025, 'Canada', 'High'), ('Jacobs and Ehman', 24, 2025, 'Not assigned', '—'),
 ('Wang et al', 7, 2025, 'Sweden', 'High'), ('Antosz-Rekucka and Prochwicz', 25, 2025, 'Poland', 'High'),
 ('Innab et al', 26, 2025, 'Saudi Arabia', 'High'), ('Okajima and Okajima', 48, 2025, 'Not assigned', '—'),
 ('Shahzad et al', 49, 2025, 'Not assigned', '—'), ('Westermark et al', 6, 2024, 'Sweden', 'High'),
 ('Lee et al', 27, 2024, 'Hong Kong, China', 'High'), ('Brown et al', None, 2024, 'Not assigned', '—'),
 ('Loukzadeh et al', 50, 2024, 'Not assigned', '—'), ('Winslow et al', 14, 2023, 'Online (Reddit)', '—'),
 ('Akyuz Cim and Cim', 28, 2023, 'Türkiye', 'Upper-middle'), ('Ekmekçi Ertek et al', 29, 2023, 'Türkiye', 'Upper-middle'),
 ('Li et al', 13, 2023, 'Taiwan', 'High'), ('Senín-Calderón et al', 30, 2023, 'Spain', 'High'),
 ('Liguori et al', 51, 2023, 'Not assigned', '—'), ('Mahmood et al', None, 2023, 'Not assigned', '—'),
 ('Kulkarni et al', 31, 2022, 'Australia', 'High'), ('Eisenlohr-Moul et al', 5, 2022, 'Online (international)', '—'),
 ('Pekçetin et al', 32, 2022, 'Türkiye', 'Upper-middle'), ('Hantsoo et al', 18, 2022, 'Online (international)', '—'),
 ('Yang et al', 33, 2022, 'Iceland', 'High'), ('Younes et al', 16, 2021, 'Lebanon', 'Lower-middle'),
 ('Beddig et al', 34, 2020, 'Germany', 'High'), ('Izadi-Mazidi and Amiri', 35, 2019, 'Iran', 'Upper-middle'),
 ('Śliwerski and Bielawska-Batorowicz', 36, 2019, 'Poland', 'High'), ('Shams-Alizadeh et al', 37, 2018, 'Iran', 'Upper-middle'),
 ('Lorenz et al', 38, 2017, 'Not assigned', '—'), ('Schmalenberger et al', 39, 2017, 'United States', 'High'),
 ('Ducasse et al', 40, 2016, 'France', 'High'), ('Petersen et al', 41, 2016, 'United States', 'High'),
 ('Balık et al', 42, 2015, 'Türkiye', 'Upper-middle'), ('Adegoke et al', 17, 2014, 'Nigeria', 'Lower-middle'),
 ('Ko et al', 43, 2013, 'Taiwan', 'High'), ('Pilver et al', 4, 2013, 'United States', 'High'),
 ('Hong et al', 44, 2012, 'South Korea', 'High'), ('Delara et al', 45, 2012, 'Iran', 'Upper-middle'),
 ('Pilver et al', 46, 2011, 'United States', 'High'), ('Heinemann et al', 47, 2010, 'Multinational', '—'),
]
assert len(TABLE_ROWS) == 44
TABLE_NUMS = [num(r[1]) if r[1] else None for r in TABLE_ROWS]
T('Publication increased sharply over the review period: 6 studies were published in 2010–2013, 6 in 2014–2017, 5 in 2018–2021 and 27 (61%) in 2022–2025 (Figure 3). Most were cross-sectional surveys of university students or community samples. The remainder included a nationwide longitudinal cohort,{13} a prospective cohort,{6} ambulatory daily-rating studies,{34} a case-control study of suicide attempters,{37} an interview-based qualitative study{19} and a thematic analysis of an online PMDD community.{14}')
H2('Psychosocial Outcomes')
T('Depression was the most frequently documented outcome (39 studies), followed by psychological distress (24), interpersonal functioning (18), suicidal ideation or self-harm (14) and quality of life (14); only 3 studies examined coping mechanisms (Figure 2).')
T('***Depression.*** In a nationwide Taiwanese cohort, women with PMDD had a 2.58-fold higher hazard of later unipolar depression and were diagnosed at a younger age than controls.{13} Depression was the most prominent feature of PMDD in a Taiwanese clinical sample,{43} and PMDD severity, rather than the diagnosis itself, was associated with sensitivity to depression in Türkiye.{28} Among Lebanese university students, depression mediated the association between childhood psychological and sexual abuse and PMDD.{16} Negative cognitive styles were more common among women with premenstrual disorders,{36} whereas a Spanish study found that self-reported PMDD (51.8%) greatly overestimated clinically confirmed PMDD (5.9%) and did not support a cognitive vulnerability to depression.{30}'.replace('***', '**'))
T('***Psychological distress.*** Women with PMDD reported greater anxiety sensitivity and difficulty regulating emotions, especially in the luteal phase,{23,41} and ambulatory studies showed stronger links between rumination and later negative mood than in controls.{34} Distress was associated with low resilience and loneliness,{27} perceived stress and poor sleep,{26} anxious and cyclothymic temperament,{29} neuroticism,{35} trait anger{40} and pain sensitivity.{25} Adversity was a recurring theme: 83% of women with PMDD in an Australian study reported early-life trauma,{31} adverse childhood experiences were associated with premenstrual disorders in a dose-dependent manner,{33} and perceived discrimination was associated with a higher likelihood of PMDD.{46}'.replace('***', '**'))
T('***Interpersonal functioning.*** People with PMDD and their partners both reported lower quality of life and relationship quality than controls.{21} In a Swedish cohort, married or cohabiting women with probable severe premenstrual disorders had a higher risk of relationship disruption.{6} Other studies documented interference with social activities and relationships,{47} reduced occupational competence among university students{32} and disruptions at work.{50}'.replace('***', '**'))
T('***Suicidal ideation and self-harm.*** In a nationally representative US sample, PMDD was associated with suicidal ideation, plans and attempts after adjustment for other psychiatric conditions.{4} A global sample of 599 patients with prospectively confirmed PMDD reported high lifetime rates of suicidal thoughts and behaviors.{5} Similar associations were reported in Bangladesh,{15} Iran{37} and South Korea,{44} and PMDD was associated with trait anger among female suicide attempters in France,{40} and 1 article argued that suicidality should be considered for inclusion in the diagnostic criteria for PMDD.{22}'.replace('***', '**'))
T('***Quality of life.*** PMDD was associated with lower quality of life than premenstrual syndrome in Sweden{7} and Türkiye{42} and with poorer health-related quality of life among Iranian adolescents.{45} Reduced quality of life was also reported in multinational{47} and physical-activity studies.{51}'.replace('***', '**'))
T('***Coping mechanisms.*** Only 3 studies examined how women cope with PMDD. An analysis of an anonymous online community found that members used shared narratives, camaraderie and information-sharing to fill gaps in social and medical support.{14} Related work pointed to psychological flexibility{48} and emotion regulation strategies{34,41} as possible targets for support, but no study evaluated coping in a low- or lower-middle-income setting. Studies of care experiences described misdiagnosis, masking of symptoms and uneven provider knowledge,{18,19} and another study reported limited awareness of PMDD.{49}'.replace('***', '**'))
H2('Geographic Distribution')
T('Thirty of the 44 studies had an identifiable country setting (Figure 4). Twenty came from high-income countries: the United States (4); Sweden, Poland and Taiwan (2 each); and Australia, Canada, France, Germany, Iceland, Saudi Arabia, South Korea, Spain, the United Kingdom and Hong Kong, China (1 each). Seven came from upper-middle-income countries (Türkiye, 4; Iran, 3). Only 3 (6.8% of 44) came from low- and lower-middle-income countries: Bangladesh,{15} Lebanon{16} and Nigeria.{17} The Nigerian study, a survey of secondary school adolescents in Ibadan published in 2014, was the only study from sub-Saharan Africa.{17} The remaining 14 studies used online or multinational samples or did not report a single setting in the available records.')

# ---------- discussion ----------
H1('Discussion')
T('This scoping review found that research on the psychosocial consequences of PMDD is growing, with most studies published since 2022, but that it is concentrated on depression and distress and on high-income countries. The focus on depression, distress and suicidality is consistent with population studies linking PMDD to these outcomes.{4–7}'.replace('4–7', '4,5,6,7'))
T('The clearest gap is coping. Only 3 studies examined how women manage PMDD, so little is known about which strategies or supports help, or how coping differs across cultures and health systems. The available evidence suggests that online peer communities fill gaps left by health services{14} and that emotion regulation and psychological flexibility may be useful targets for support.{34,41,48} Because PMDD symptoms recur every month, coping is not a secondary question: it shapes how women function at home, at school and at work between clinical encounters.')
T('The second gap is geographic. Premenstrual symptoms are common across Africa,{8} yet only 1 study, published more than a decade ago, came from sub-Saharan Africa.{17} This pattern most likely reflects limited research rather than low burden. Stigma around menstruation, limited mental health services and local explanations of menstrual distress may shape both the burden of PMDD and how women cope in these settings, but these questions have not been studied. The low- and lower-middle-income studies that do exist sampled university or school students,{15–17} so evidence from community, rural and clinical populations is absent.'.replace('15–17', '15,16,17'))
T('Diagnostic practice also limits what this evidence can show. Much PMDD research relies on provisional diagnosis from retrospective questionnaires, which can overestimate prevalence and severity.{1,30} Daily mood varies substantially within and between cycles,{38} and the number and type of symptoms that predict impairment are still debated.{39} Future studies should confirm PMDD with prospective daily ratings{1,20} and pair confirmed diagnosis with measures of coping and help-seeking.')
H2('Limitations')
T('This review has several limitations. First, studies outside the 5 databases or the 2010–2025 period may have been missed. Second, eligibility was assessed from titles, abstracts and subject terms without full-text review, so some included records, such as a commentary, a narrative review and studies of premenstrual disorders more broadly, may not meet the criteria on full-text review. Third, studies varied in design and in how outcomes and PMDD were defined, which limits direct comparison, and a study could be coded to more than 1 domain. Fourth, a country setting could not be assigned to 14 studies, and complete bibliographic details could not be located for 2 included records. Finally, as is standard for scoping reviews, we did not appraise study quality.{9,11}')
H2('Public Health Implications')
T('Our findings point to 4 practical steps. First, clinicians should ask about suicidal thoughts when women present with severe premenstrual mood symptoms.{4,5,15} Second, PMDD screening can be integrated into sexual and reproductive health, school and university health services, where many affected young women already seek care. Third, provider training can improve recognition and reduce misdiagnosis and delayed care.{18,19} Fourth, services should ask how women manage their symptoms and connect them with peer and professional support.{14} For research, priorities include prospective confirmation of PMDD, qualitative studies of stigma and coping in African and other low-income settings, validation of screening tools in local languages and evaluation of coping-focused interventions.')
H2('Conclusions')
T('PMDD carries a well-documented psychosocial burden, especially depression, distress and suicidality, yet how women cope has received little attention, and evidence from low-income settings and Africa is nearly absent. Culturally responsive screening, integrated mental and reproductive health services and more research in African settings are needed so that health systems can recognize and respond to this burden.')

# ---------- statements ----------
page_break()
H1('Contributors')
P('[[Describe each author’s contribution (e.g., I. P. Shobayo conceptualized the study and led screening and writing; other authors contributed to screening, data charting, interpretation and revision). All authors approved the final version.]]', indent=False)
H1('Acknowledgments')
P('None.', indent=False)
H1('Conflicts of Interest')
P('The authors have no conflicts of interest to disclose. [[Confirm with all co-authors.]]', indent=False)
H1('Human Participant Protection')
P('No protocol approval was needed because this study used only published data.', indent=False)

# ---------- references ----------
page_break()
H1('References')
for i, k in enumerate(order, 1):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.35); p.paragraph_format.first_line_indent = Inches(-0.35)
    p.add_run(f'{i}.\t{POOL[k]}')
unused = sorted(set(POOL) - set(order))
assert not unused, unused

# ---------- figure legends ----------
page_break()
H1('Figure Legends')
LEG = [
 ('Figure 1.', 'Study Selection Flow Diagram', ' Adapted from PRISMA-ScR.{11}'),
 ('Figure 2.', 'Psychosocial Outcomes Reported in the 44 Included Studies', ' A study could report more than 1 outcome, so counts should not be summed.'),
 ('Figure 3.', 'Included Studies by Publication Period', ' Twenty-seven of 44 studies (61%) were published in 2022–2025.'),
 ('Figure 4.', 'Country Setting and World Bank Income Group of Included Studies', ' The map shows the 30 studies with an identifiable country setting; income groups follow the World Bank FY2027 classification.{12} The other 14 studies used online or multinational samples or had no single setting in the available records. The high-income count includes 1 study from Hong Kong, China.'),
]
for a, b, c in LEG: P(f'**{a} {b}.**{c}', indent=False)

# ---------- table 1 ----------
page_break()
P('**Table 1. Included Studies by Publication Year, Country Setting and World Bank Income Group (n = 44)**', indent=False, single=True)
tb = doc.add_table(rows=1, cols=4); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, h in zip(tb.rows[0].cells, ['Study', 'Year', 'Country setting', 'Income group']):
    c.text = ''; rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(10)
for (name, ref, yr, ctry, inc), n in zip(TABLE_ROWS, TABLE_NUMS):
    cells = tb.add_row().cells
    for c in cells: c.paragraphs[0].paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p0 = cells[0].paragraphs[0]; rr = p0.add_run(name); rr.font.size = Pt(10)
    if n:
        rr = p0.add_run(str(n)); rr.font.superscript = True; rr.font.size = Pt(10)
    else:
        rr = p0.add_run(' [complete citation not located]'); rr.font.size = Pt(10); rr.bold = True; rr.font.color.rgb = NOTE
    for c, v in zip(cells[1:], [str(yr), ctry, inc]):
        rr = c.paragraphs[0].add_run(v); rr.font.size = Pt(10)
for row in tb.rows:
    for c, w in zip(row.cells, [2.9, 0.6, 1.9, 1.1]): c.width = Inches(w)
P('*Note.* Income groups follow the World Bank FY2027 classification.{12} “Not assigned” indicates that a single country setting could not be determined from the available records.', indent=False, size=10, single=True)

# ---------- figures ----------
for f, (a, b, c) in zip(['figure1_prisma.png', 'figure2_outcomes.png', 'figure3_period.png', 'figure4_map.png'], LEG):
    page_break()
    doc.add_picture(f'/root/work/pmdd/ms/{f}', width=Inches(6.2 if 'map' in f else 5.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    P(f'**{a} {b}.**{c}', indent=False, single=True)

# ---------- word counts ----------
main_words = sum(len(re.sub(r'\{[\d,\s]+\}|\[\[.+?\]\]|\*', '', t).split()) for t in WORDS)
runs(wc_par, f'**Word count:** abstract {abs_words} words; main text {main_words} words; 1 table; 4 figures; {len(order)} references.')
doc.save('/root/work/pmdd/ms/PMDD_Scoping_Review_Manuscript.docx')
print('abstract', abs_words, 'main', main_words, 'refs', len(order))
