import pandas as pd, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
log=pd.read_pickle('log.pkl'); T=pd.read_pickle('table.pkl'); S=json.load(open('stats.json'))
wb=Workbook()
NAVY='0A254E'; hdr_font=Font(bold=True,color='FFFFFF',name='Arial',size=10); hdr_fill=PatternFill('solid',fgColor=NAVY)
body=Font(name='Arial',size=10); thin=Side(style='thin',color='D0D5DD')
def sheet(ws,df,widths,wrap_cols=()):
    ws.append(list(df.columns))
    for r in df.itertuples(index=False): ws.append(list(r))
    for c in ws[1]: c.font=hdr_font; c.fill=hdr_fill; c.alignment=Alignment(wrap_text=True,vertical='center')
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=body; c.alignment=Alignment(wrap_text=(c.column-1) in wrap_cols,vertical='top')
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
# README
ws=wb.active; ws.title='README'
lines=[('PMDD Psychosocial Outcomes & Coping – Scoping Review: Screening and Charting Workbook',True),
('',False),
('What this workbook contains',True),
('PRISMA-ScR counts – the numbers behind the flow diagram, computed from the uploaded search exports.',False),
('Screening log – every one of the 1,801 exported records, with its deduplication/screening decision and the reason for exclusion.',False),
('Included studies – the 85 included studies with charted country, World Bank income group, design, sample, PMDD ascertainment and psychosocial domains.',False),
('Domain summary / Geography summary – counts used in the manuscript and poster.',False),
('',False),
('Sources screened',True),
('EBSCOhost exports (CINAHL Ultimate, APA PsycInfo, APA PsycArticles, MEDLINE Ultimate, Women\'s Studies International), Cochrane Library export, PubMed export (6 records).',False),
('',False),
('Eligibility criteria applied',True),
('Include: peer-reviewed primary empirical (quantitative, qualitative or mixed-methods) study; published 2010–2025; English; PMDD (DSM-IV/DSM-5/ICD-defined, provisional or confirmed, incl. premenstrual disorders with a distinct PMDD group or diagnosis) as a defined study population, diagnostic group or exposure; at least one psychosocial outcome or coping mechanism as a primary aim.',False),
('Exclude: animal studies; laboratory/biological studies (neuroimaging, hormonal, genetic, physiological); intervention/treatment trials; reviews, meta-analyses, guidelines, commentaries; instrument-validation studies; case reports; books, chapters, dissertations, letters, editorials, conference abstracts; PMS-only populations.',False),
('',False),
('Important notes',True),
('Eligibility was judged on titles, abstracts and subject terms (single reviewer). Full texts were not retrieved; a second reviewer should verify decisions and full-text eligibility before final publication.',False),
('Deduplication used DOI and normalised title matching (Unicode-aware), with manual review of near-duplicate titles.',False),
('Country is taken from the abstract; where the abstract did not name it, it was taken from the authors\' institutional setting (verify at full text). One study (record 208) does not report country in its abstract.',False),
('Income groups use the World Bank FY2027 classification (effective 1 July 2026). "LMIC" in the World Bank sense = low + lower-middle + upper-middle income.',False),
('Domain codes: DEP depression; ANX anxiety; MER mood & emotion regulation; PD psychological distress & perceived stress; IF interpersonal & relational functioning; SUI suicidality & self-harm; QOL quality of life; COP coping & resilience; FUN academic/occupational functioning; TRA trauma, adversity & discrimination; COM psychiatric comorbidity; PER personality & temperament; EXP lived experience, stigma & help-seeking.',False)]
for t,b in lines:
    ws.append([t]); c=ws.cell(row=ws.max_row,column=1); c.font=Font(name='Arial',size=12 if b else 10,bold=b,color=NAVY if b else '000000'); c.alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['A'].width=140
# PRISMA
ws=wb.create_sheet('PRISMA-ScR counts'); P=S['prisma']
grp={'Published outside 2010–2025':['Published outside 2010–2025'],
 'Non-English full text':['Non-English full text'],
 'Reviews, meta-analyses, guidelines, commentaries':['Review, meta-analysis or guideline (not a primary empirical study)','Review, commentary, editorial, guideline or other non-empirical article'],
 'Other ineligible publication types (books, chapters, dissertations, letters, editorials, case reports, conference abstracts)':['Ineligible publication type (book, chapter, dissertation, letter, editorial, comment, case report, non-peer-reviewed)','Case report or case series','Conference abstract'],
 'PMDD not the primary study population':['PMDD not studied (not mentioned in title, abstract or subject terms)','PMDD not the primary/defined study population'],
 'Intervention or treatment study':['Intervention or treatment study'],
 'Laboratory/biological study':['Laboratory/biological study (neuroimaging, hormonal, genetic, physiological)'],
 'Instrument development/validation study':['Instrument development/validation study'],
 'Animal study':['Animal study'],
 'No psychosocial outcome or coping mechanism as a primary aim':['No psychosocial outcome or coping mechanism as a primary aim']}
rows=[('IDENTIFICATION','',''),('Records identified from database exports',P['identified_total'],'')]
rows+=[('   '+k,v,'by database') for k,v in P['identified_by_database'].items()]
rows+=[('Duplicate records removed',P['duplicates_removed'],''),('SCREENING','',''),('Records screened (title/abstract/subject terms)',P['screened'],''),('Records excluded',P['excluded_total'],'')]
G={k:sum(P['excluded_by_reason'].get(x,0) for x in v) for k,v in grp.items()}
assert sum(G.values())==P['excluded_total']
rows+=[('   '+k,v,'') for k,v in sorted(G.items(),key=lambda x:-x[1])]
rows+=[('INCLUDED','',''),('Studies included in the scoping review',P['included'],'')]
sheet(ws,pd.DataFrame(rows,columns=['Stage','n','Note']),[95,10,15])
json.dump(G,open('prisma_groups.json','w'),indent=1)
# screening log
ws=wb.create_sheet('Screening log')
L=log.rename(columns={'record_id':'Record ID','source_export':'Source export','database':'Database','accession':'Accession/ID','authors':'Authors','year':'Year','title':'Title','journal':'Journal','doi':'DOI','language':'Language','stage':'Stage','decision':'Decision','reason':'Exclusion reason','link':'Link'})
L['Title']=L['Title'].str.replace(r'<[^>]+>','',regex=True)
sheet(ws,L,[9,30,22,16,30,7,60,30,26,10,26,16,45,30],wrap_cols=(6,12))
# included
ws=wb.create_sheet('Included studies')
I=T.drop(columns=['domain_codes']).rename(columns={'record_id':'Record ID','citation':'Citation','year':'Year','title':'Title','journal':'Journal','doi':'DOI','country':'Country','region':'Region','income_group':'World Bank income group (FY2027)','design':'Design','design_group':'Design group','sample_size':'Sample size (N)','population':'Population','pmdd_ascertainment':'PMDD ascertainment','ascertainment_group':'Ascertainment group','domains':'Psychosocial domains'})
sheet(ws,I,[9,22,7,60,28,26,18,22,22,32,24,12,32,32,30,55],wrap_cols=(3,9,12,13,15))
# summaries
ws=wb.create_sheet('Domain summary')
D=pd.DataFrame([(k,v,round(100*v/85,1)) for k,v in sorted(S['domains'].items(),key=lambda x:-x[1])],columns=['Psychosocial domain','Studies (n)','% of 85'])
sheet(ws,D,[50,12,10])
ws=wb.create_sheet('Geography summary')
rows=[('By World Bank income group','','')]+[(k,v,round(100*v/85,1)) for k,v in S['income'].items()]+[('','',''),('By region','','')]+[(k,v,round(100*v/85,1)) for k,v in S['region'].items()]+[('','',''),('By country','','')]+[(k,v,round(100*v/85,1)) for k,v in S['country'].items()]
sheet(ws,pd.DataFrame(rows,columns=['Group','Studies (n)','% of 85']),[50,12,10])
wb.save('PMDD_Scoping_Review_Screening_and_Charting.xlsx'); print('saved', G)
