"""Draft text for a retrospective OSF registration (Generalized Systematic Review Registration
template), built from the review's actual methods and scripts/data.json.
Usage: python3 scripts/make_osf_registration.py"""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

HERE = Path(__file__).parent
ROOT = HERE.parent
d = json.loads((HERE / "data.json").read_text())
P, N, FT = d["prisma"], d["total_included"], d["full_text_checked"]

doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = "Times New Roman", Pt(11)


def text(par, s):
    """Write s into par; text inside [[ ]] is highlighted for the authors to complete."""
    for i, piece in enumerate(s.split("[[")):
        if i == 0:
            par.add_run(piece)
            continue
        todo, rest = piece.split("]]", 1)
        r = par.add_run("[" + todo + "]")
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        par.add_run(rest)


def field(name, *paras):
    doc.add_heading(name, 2)
    for p in paras:
        text(doc.add_paragraph(), p)


doc.add_heading("OSF Registration Draft (Retrospective) — Scoping Review", 1)
text(doc.add_paragraph(), "Template: OSF “Generalized Systematic Review Registration.” Paste each "
     "answer into the matching field. Highlighted items need the authors' input. This review "
     "was completed before registration, so the registration must be labelled retrospective.")

field("Title",
      "Mental Health and Psychosocial Outcomes Among Individuals With "
      "Mayer-Rokitansky-Küster-Hauser (MRKH) Syndrome: A Scoping Review")
field("Description",
      "Retrospective registration of a completed scoping review mapping the mental health and "
      "psychosocial outcomes, coping mechanisms, and healthcare system gaps reported in primary "
      "studies of individuals with MRKH syndrome published January 2019–March 2026. This "
      "registration describes the methods as conducted; it was not written before the review "
      "began.")
field("Contributors",
      "Ifeoluwanimi P. Shobayo, BSc, MSPHc (Department of Public and Community Health, Liberty "
      "University); Paul Okojie, PhD (Liberty University); Robyn Anderson, PhD (Liberty "
      "University).")
field("Tasks and roles",
      "[[Confirm: I.P.S.—conceptualization, search, screening, data charting, analysis, "
      "writing; P.O. and R.A.—supervision, methodology, review and editing.]]")

doc.add_heading("Type of review and stages", 1)
field("Type of review", "Scoping review.")
field("Review stages",
      "Preparation; search; title/abstract screening; eligibility (full-text) assessment; data "
      "charting; synthesis; reporting.")
field("Current review stage",
      "Completed. All stages were finished before registration (retrospective registration).")
field("Start date", "[[Month year the review was started]]")
field("End date", "September 2026 (data charting and synthesis completed).")

doc.add_heading("Background and questions", 1)
field("Background",
      "MRKH syndrome is the congenital absence of the uterus and upper vagina in individuals "
      "with a 46,XX karyotype and functioning ovaries, affecting about 1 in 4,500–5,000 female "
      "births. It is usually diagnosed in adolescence and affects fertility, sexuality, and "
      "identity. Earlier studies linked MRKH with psychological distress, anxiety, and reduced "
      "quality of life, and prior reviews focused on psychological and sexual outcomes. Coping "
      "mechanisms and healthcare system gaps have not been mapped across the literature "
      "published since 2019.")
field("Primary research question",
      "What mental health and psychosocial outcomes and coping mechanisms are reported among "
      "individuals with MRKH syndrome, and what healthcare system gaps exist?")
field("Secondary research questions",
      "What study designs and geographic regions make up the recent evidence base, and where are "
      "the gaps?")
field("Expectations / hypotheses",
      "None; scoping reviews map evidence and do not test hypotheses.")

doc.add_heading("Framework and eligibility", 1)
field("Methodological framework",
      "Arksey and O'Malley's five-stage framework (Int J Soc Res Methodol. 2005;8(1):19-32); "
      "reported according to PRISMA-ScR (Ann Intern Med. 2018;169(7):467-473).")
field("Population, Concept, Context",
      "Population: individuals with MRKH syndrome (Müllerian agenesis/aplasia). Concept: mental "
      "health and psychosocial outcomes, coping mechanisms, and healthcare system gaps. Context: "
      "any country or care setting.")
field("Inclusion criteria",
      "(1) Primary quantitative, qualitative, or mixed-methods research; (2) published January "
      "2019–March 2026; (3) English language; (4) participants with MRKH; mixed samples only if "
      "results were reported separately for participants with MRKH or at least 80% of "
      "participants had MRKH; (5) reports at least one psychological, mental health, or "
      "psychosocial outcome, or coping.")
field("Exclusion criteria (codes used at screening)",
      "E1 published before 2019; E2 not primary research (review, commentary, letter, erratum, "
      "chapter); E3 case report, small surgical case series, or surgical video; E4 wrong "
      "population or unrelated topic, including mixed samples without separate MRKH results; E5 "
      "anatomical, surgical, functional, or fertility outcomes only; E6 not published in English.")

doc.add_heading("Search", 1)
field("Databases and interfaces",
      "PubMed (MEDLINE); Scopus; EBSCOhost: MEDLINE Ultimate, CINAHL Ultimate, APA PsycInfo, and "
      "Women's Studies International.")
field("Search dates",
      "Records exported March 29, 2026. [[Confirm the PubMed search was run on the same date.]]")
field("Records retrieved",
      f"{P['identified']} in total: EBSCOhost combined export 207 (MEDLINE 145, CINAHL 39, Women's "
      "Studies International 15, APA PsycInfo 8); APA PsycInfo separate export 20; PubMed 146; "
      "Scopus 143.")
field("Query strings",
      "Searches combined terms for MRKH (e.g., Mayer-Rokitansky-Küster-Hauser, MRKH, Müllerian "
      "agenesis, vaginal agenesis) with terms for mental health, psychosocial outcomes, and "
      "coping (e.g., depression, anxiety, psychological distress, quality of life, body image, "
      "coping). [[Paste the exact search string for each database from your search histories.]]")
field("Date limits",
      "Date limits were not applied identically across interfaces (Scopus 2019–2026; PubMed "
      "export 2021–2026, with 2019–2020 MEDLINE records captured through EBSCOhost; EBSCOhost "
      "exported without a date limit). Records published before 2019 were removed at "
      "title/abstract screening (criterion E1).")
field("Grey literature and other sources",
      "Grey literature was not searched. [[State whether reference lists of included studies "
      "were hand-searched; if not, write “None.”]]")
field("Contacting authors", "Study authors were not contacted.")
field("Search repetition", "The search was not repeated after March 29, 2026.")

doc.add_heading("Screening", 1)
field("Deduplication",
      f"Records from all exports were merged and duplicates identified by matching DOIs and "
      f"normalized titles: {P['duplicates_removed']} duplicates removed, leaving "
      f"{P['screened']} unique records.")
field("Screening stages",
      f"Stage 1: titles and abstracts ({P['screened']} records; {P['excluded']} excluded). "
      f"Stage 2: eligibility assessment of {P['fulltext']} records ({P['excluded_eligibility']} "
      f"excluded: mixed sample without separate MRKH results, 5; no psychological outcome, 2), "
      f"leaving {N} included studies. Full texts were obtained for {FT} of the {N} included "
      "studies; the remaining 5 were assessed on the published abstract or title.")
field("Screening procedure and software",
      "Initial screening (March 2026) was conducted in Rayyan with references managed in "
      "Zotero. Because the original screening records could not be recovered, screening was "
      "repeated in September 2026 from the original export files. In the repeated screening, "
      "an AI tool (Claude, Anthropic) proposed a decision and exclusion code for every record, "
      "and each decision was recorded in a screening workbook. [[Describe human verification: "
      "e.g., “The first author and a second reviewer independently verified every decision; "
      "disagreements were resolved by discussion.” Only state what was actually done.]]")
field("Screening reliability and reconciliation",
      "[[Describe how agreement between reviewers was checked and how disagreements were "
      "resolved, or state that single-reviewer verification was used.]]")

doc.add_heading("Data charting (extraction)", 1)
field("Entities extracted",
      "Author; year; country and world region; study design (quantitative, qualitative, mixed "
      "methods) and design detail; sample size; outcome domains; coping mechanisms; healthcare "
      "system gaps; verification notes.")
field("Coding scheme",
      "Outcome domains (a study could have several): depression and anxiety; quality of life, "
      "body image, and self-esteem; psychosexual and relational challenges; broader psychological "
      "distress; coping mechanisms documented; healthcare system gaps. Coping categories: peer, "
      "family, and social support; psychological counseling or intervention; identity "
      "reconstruction or positive reappraisal; spiritual coping; acceptance; avoidance or "
      "concealment; self-advocacy.")
field("Charting procedure",
      f"Data were charted from the full text for {FT} of {N} studies and from the abstract or "
      "title for the remaining 5. Charting was AI-assisted (Claude, Anthropic) and checked "
      "against each full-text PDF. [[Describe any second-reviewer check of charted data.]]")

doc.add_heading("Quality assessment and synthesis", 1)
field("Critical appraisal",
      "Not performed, consistent with scoping review methodology.")
field("Synthesis",
      "Descriptive numerical summary (number and proportion of included studies reporting each "
      "outcome domain, coping category, study design, and region) and qualitative content "
      "analysis grouping findings into outcome domains, coping mechanisms, and healthcare system "
      "gaps. Counts describe how often outcomes were studied, not their prevalence.")
field("Handling of overlapping samples",
      "Reports that appeared to share samples (two Chinese reports; two Danish reports) were "
      "retained as separate reports and flagged.")

doc.add_heading("Other", 1)
field("Funding", "No external funding.")
field("Conflicts of interest",
      "[[The authors declare no conflicts of interest — confirm with all authors.]]")
field("Data and materials",
      "The screening workbook (decision and reason for every record) and the data-charting table "
      "will be uploaded to this OSF project. [[Upload: MRKH_rescreening_workbook.xlsx and "
      "extraction_provisional.csv (rename, e.g., “MRKH_data_charting.csv”).]]")

out = ROOT / "osf" / "OSF_Registration_Draft.docx"
doc.save(out)
print("wrote", out)
