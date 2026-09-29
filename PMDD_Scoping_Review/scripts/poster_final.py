import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree
S=json.load(open('stats.json')); G=json.load(open('prisma_groups.json')); P=S['prisma']
NAVY=RGBColor(0x0A,0x25,0x4E); WHITE=RGBColor(255,255,255); INK=RGBColor(0x11,0x18,0x27); MUTED=RGBColor(0x4B,0x55,0x63)
RED=RGBColor(0xA3,0x1F,0x34); GRAY=RGBColor(0x9C,0xA3,0xAF); LIGHT=RGBColor(0xF3,0xF4,0xF6); PALE=RGBColor(0xEE,0xF2,0xF8); GREEN=RGBColor(0x04,0x78,0x57); DIV=RGBColor(0xD1,0xD5,0xDB)
FONT='Times New Roman'
prs=Presentation('/root/work/prev.pptx'); s=prs.slides[0]
keep={'TextBox 3','Picture 76','TextBox 304','Picture 18'}
for sh in list(s.shapes):
    if sh.name not in keep: sh._element.getparent().remove(sh._element)
# clear old title text (keep its frame)
for sh in s.shapes:
    if sh.name=='TextBox 3':
        tf=sh.text_frame
        for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
        for r in list(tf.paragraphs[0].runs): r._r.getparent().remove(r._r)
def tb(x,y,w,h,fill=None,line=None,anchor=MSO_ANCHOR.TOP,margin=0.03,shape=MSO_SHAPE.RECTANGLE,lw=0.6):
    if fill is None and line is None and shape==MSO_SHAPE.RECTANGLE: b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    else:
        b=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h)); b.shadow.inherit=False
        if fill is None: b.fill.background()
        else: b.fill.solid(); b.fill.fore_color.rgb=fill
        if line is None: b.line.fill.background()
        else: b.line.color.rgb=line; b.line.width=Pt(lw)
        if shape==MSO_SHAPE.ROUNDED_RECTANGLE: b.adjustments[0]=0.08
    tf=b.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(0.015)
    return b
def write(b,paras,size=4.4,color=INK,align=PP_ALIGN.LEFT,space=1.2,ls=None):
    tf=b.text_frame; first=True
    for it in paras:
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        if isinstance(it,dict): runs=it['runs']; bullet=it.get('bullet',False); sz=it.get('size',size); al=it.get('align',align); col=it.get('color',color); sp=it.get('space',space)
        else: runs=it; bullet=False; sz=size; al=align; col=color; sp=space
        if isinstance(runs,str): runs=[(runs,False,False)]
        p.alignment=al; p.space_after=Pt(sp)
        if ls: p.line_spacing=ls
        if bullet:
            pPr=p._p.get_or_add_pPr(); pPr.set('marL',str(int(Inches(0.08)))); pPr.set('indent',str(-int(Inches(0.07))))
            bu=etree.SubElement(pPr,qn('a:buFont')); bu.set('typeface','Arial'); bc=etree.SubElement(pPr,qn('a:buChar')); bc.set('char','•')
        for r_ in runs:
            t,bo,itl=r_[:3]; c=r_[3] if len(r_)>3 else col
            r=p.add_run(); r.text=t; f=r.font; f.name=FONT; f.size=Pt(sz); f.bold=bo; f.italic=itl; f.color.rgb=c
    return b
def header(x,y,w,h,text,size=10):
    b=tb(x,y,w,h,fill=NAVY,anchor=MSO_ANCHOR.MIDDLE); write(b,[text],size=size,color=WHITE,align=PP_ALIGN.CENTER,space=0); return b
def body(x,y,w,h,paras,size=4.4,**k):
    b=tb(x,y,w,h,fill=WHITE,margin=0.05); b.text_frame.margin_top=Inches(0.04); write(b,paras,size=size,**k); return b
# ---------------- TITLE ----------------
t=tb(1.4,0.075,8.45,0.675,anchor=MSO_ANCHOR.MIDDLE)
write(t,[{'runs':[('PSYCHOSOCIAL OUTCOMES AND COPING MECHANISMS AMONG WOMEN WITH PREMENSTRUAL DYSPHORIC DISORDER: A GLOBAL SCOPING REVIEW',True,False)],'size':10,'align':PP_ALIGN.CENTER,'space':1.5},
         {'runs':[('Ifeoluwanimi P. Shobayo, MSPH, Chelsea R. Mazonde, MPH, Marylyn O. Oduneye, MSPH, Tahirou Diallo, MSPH, Fadzai G. Nyarugwe, MSPH, Cynthia C. Ilechukwu, MSPH',False,False)],'size':7.9,'align':PP_ALIGN.CENTER,'space':0}])
# ---------------- LEFT COLUMN ----------------
LX,LW=0.107,2.217
header(LX,0.806,LW,0.199,'Abstract')
body(LX,1.03,LW,1.52,[
 [('Background: ',True,False),('Premenstrual dysphoric disorder (PMDD) is a DSM-5 depressive disorder affecting an estimated 1.6% of women and girls (~31 million) worldwide.¹ This scoping review mapped psychosocial outcomes and coping mechanisms among women with PMDD and the geographic distribution of the evidence.',False,False)],
 [('Methods: ',True,False),("Following Arksey and O’Malley’s framework and PRISMA-ScR, records from CINAHL, APA PsycInfo, APA PsycArticles, MEDLINE, Women’s Studies International, PubMed and the Cochrane Library (2010–2025) were screened. Eligible peer-reviewed empirical or qualitative studies had PMDD as the population and a psychosocial outcome or coping mechanism as a primary aim.",False,False)],
 [('Results: ',True,False),(f"Of {P['identified_total']:,} records, {P['duplicates_removed']} were duplicates; {P['screened']:,} were screened and {P['included']} studies were included. Most studied were depression (n=28), mood and emotion regulation (n=18), psychiatric comorbidity (n=18), psychological distress (n=16) and suicidal ideation or self-harm (n=13); only 7 studies examined coping. Two-thirds (67%) came from high-income countries; 3 (3.5%) came from sub-Saharan Africa (Ethiopia, n=2; Nigeria, n=1).",False,False)],
 [('Conclusions: ',True,False),('Evidence on the psychosocial burden of PMDD is growing but is concentrated in high-income settings, relies on provisional diagnoses and rarely addresses coping.',False,False)],
 [('Keywords: ',True,False),('premenstrual dysphoric disorder; psychosocial outcomes; coping mechanisms; LMICs; scoping review',False,True)]],size=4.25,align=PP_ALIGN.JUSTIFY,space=1.2)
header(LX,2.62,LW,0.36,'Introduction, Objective, and Research Question',size=9.5)
body(LX,3.03,LW,1.36,[
 [('Introduction: ',True,False),('PMDD is a recurring, hormone-linked mood disorder in which irritability, mood swings, depressed mood and anxiety emerge in the late luteal phase and remit after menses begin; DSM-5 classifies it as a depressive disorder.² Using prospectively confirmed diagnoses, community prevalence is about 1.6%, roughly 31 million women and girls worldwide.¹ PMDD has been linked with suicidal thoughts, strained relationships and reduced quality of life, but this evidence is scattered across disciplines, and it is unclear how women cope or where the research has been conducted.',False,False)],
 [('Objective: ',True,False),('To map psychosocial outcomes and coping mechanisms reported among women with PMDD and identify geographic gaps in the published literature.',False,False)],
 [('Research Question: ',True,False),('What psychosocial outcomes and coping mechanisms have been reported in studies of women with PMDD published between 2010 and 2025, and how are these studies distributed across geographic settings?',False,False)]],size=4.6,align=PP_ALIGN.JUSTIFY,space=1.8)
header(LX,4.46,LW,0.199,'Methods')
rows=[('Study Design','Scoping review guided by Arksey and O’Malley’s framework³ and reported using PRISMA-ScR.⁴'),
('Databases','CINAHL Ultimate, APA PsycInfo, APA PsycArticles, MEDLINE Ultimate, Women’s Studies International, PubMed/MEDLINE and Cochrane Library.'),
('Publication Years','2010–2025; English language.'),
('Inclusion Criteria','Peer-reviewed empirical or qualitative studies with PMDD (provisional or confirmed) as the study population and ≥1 psychosocial outcome or coping mechanism as a primary aim.'),
('Exclusion Criteria','Animal, laboratory/biological, intervention, review, instrument-validation and case-report studies; non-peer-reviewed publications; PMS-only samples.'),
('Screening Process',f"{P['identified_total']:,} records deduplicated (DOI and title matching) to {P['screened']:,}; titles, abstracts and subject terms screened with every exclusion reason logged; {P['included']} studies included."),
('Evidence Mapping','Studies charted by psychosocial domain (13 categories), coping mechanism, PMDD diagnostic method, country, World Bank region and income group (FY2027).')]
top=4.69; H=7.36-top
gt=s.shapes.add_table(len(rows)+1,2,Inches(LX),Inches(top),Inches(LW),Inches(H)).table
gt.columns[0].width=Inches(0.62); gt.columns[1].width=Inches(LW-0.62)
gt.rows[0].height=Inches(0.14)
for i in range(1,len(rows)+1): gt.rows[i].height=Inches((H-0.14)/len(rows))
def cell(c,txt,bold=False,fill=WHITE,size=4.1):
    c.fill.solid(); c.fill.fore_color.rgb=fill; c.vertical_anchor=MSO_ANCHOR.MIDDLE
    c.margin_left=c.margin_right=Inches(0.03); c.margin_top=c.margin_bottom=Inches(0.01)
    p=c.text_frame.paragraphs[0]; p.text=''; r=p.add_run(); r.text=txt; r.font.size=Pt(size); r.font.name=FONT; r.font.bold=bold; r.font.color.rgb=INK
for j,h in enumerate(['Category','Description']): cell(gt.cell(0,j),h,True,RGBColor(0xD9,0xE2,0xF3),4.3)
for i,(a,b_) in enumerate(rows,1):
    cell(gt.cell(i,0),a,True,PALE if i%2 else WHITE); cell(gt.cell(i,1),b_,False,PALE if i%2 else WHITE)
tblPr=gt._tbl.tblPr; tblPr.set('bandRow','0')
# ---------------- CENTER PANEL ----------------
CX,CY,CW=2.536,0.794,5.179
# stat row
tiles=[(f"{P['identified_total']:,}",'records identified','across 7 databases',NAVY),(str(P['included']),'studies included','in the scoping review',NAVY),
       ('7','studies addressing coping','8.2% of 85',RED),('3','studies from sub-Saharan Africa','3.5% of 85',RED)]
tw=(CW-0.2)/4
for k,(big,lab,sub,col) in enumerate(tiles):
    x=CX+0.1+k*tw
    b=tb(x,CY+0.07,tw,0.86,anchor=MSO_ANCHOR.MIDDLE)
    write(b,[{'runs':[(big,True,False,col)],'size':24,'align':PP_ALIGN.CENTER,'space':0},{'runs':[(lab,False,False)],'size':6,'align':PP_ALIGN.CENTER,'space':0},{'runs':[(sub,False,True,MUTED)],'size':5.5,'align':PP_ALIGN.CENTER,'space':0}])
    if k:
        ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(CY+0.15),Inches(x),Inches(CY+0.85)); ln.line.color.rgb=DIV; ln.line.width=Pt(0.5)
# --- Fig 1 PRISMA ---
fx,fy=2.64,1.78
write(tb(fx,fy,2.45,0.16),[[('Fig 1. Study Selection — PRISMA-ScR Flow Diagram',True,False)]],size=5.2,color=NAVY)
def pbox(x,y,w,h,lines,fill=PALE,line=NAVY,size=3.6,bold_first=True,align=PP_ALIGN.CENTER):
    b=tb(x,y,w,h,fill=fill,line=line,anchor=MSO_ANCHOR.MIDDLE,shape=MSO_SHAPE.ROUNDED_RECTANGLE,margin=0.03)
    write(b,[{'runs':[(l,bold_first and k==0,False)],'align':align,'space':0} for k,l in enumerate(lines)],size=size); return b
def arrow(x1,y1,x2,y2):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); c.line.color.rgb=NAVY; c.line.width=Pt(0.9)
    ln=c.line._get_or_add_ln(); te=etree.SubElement(ln,qn('a:tailEnd')); te.set('type','triangle'); te.set('w','med'); te.set('len','med')
def stage(y,h,label):
    b=tb(fx,y,0.33,h,fill=WHITE,line=GRAY,anchor=MSO_ANCHOR.MIDDLE,shape=MSO_SHAPE.ROUNDED_RECTANGLE,margin=0.01,lw=0.4)
    write(b,[{'runs':[(label,True,True)],'align':PP_ALIGN.CENTER}],size=3.4,color=NAVY)
bx,bw=3.02,1.05; rx,rw=4.14,0.97
rowsY=[(2.0,0.46,'Identification'),(2.56,0.36,'De-duplication'),(3.02,0.42,'Screening'),(3.56,0.42,'Included')]
for y,h,l in rowsY: stage(y,h,l)
db=P['identified_by_database']
pbox(bx,2.0,bw,0.46,['Records identified via','database search',f"n = {P['identified_total']:,}"],size=3.7)
pbox(rx,1.98,rw,0.5,[f"{k}: {v:,}" for k,v in db.items()],fill=WHITE,line=GRAY,size=2.9,bold_first=False,align=PP_ALIGN.LEFT)
arrow(bx+bw/2,2.46,bx+bw/2,2.56)
pbox(bx,2.56,bw,0.36,['Records after duplicate removal',f"{P['duplicates_removed']} duplicates removed",f"n = {P['screened']:,}"],size=3.6)
arrow(bx+bw/2,2.92,bx+bw/2,3.02)
pbox(bx,3.02,bw,0.42,['Records screened','(title, abstract, subject terms)',f"n = {P['screened']:,}"],size=3.6)
short={'PMDD not the primary study population':'PMDD not primary population','Published outside 2010–2025':'Outside 2010–2025','Reviews, meta-analyses, guidelines, commentaries':'Reviews / commentaries','Other ineligible publication types (books, chapters, dissertations, letters, editorials, case reports, conference abstracts)':'Other ineligible formats','Non-English full text':'Non-English','Intervention or treatment study':'Intervention studies','Laboratory/biological study':'Laboratory / biological','No psychosocial outcome or coping mechanism as a primary aim':'No psychosocial/coping aim','Instrument development/validation study':'Instrument validation','Animal study':'Animal studies'}
ex=sorted(G.items(),key=lambda x:-x[1])
pbox(rx,2.62,rw,0.98,[f"Excluded (n = {P['excluded_total']:,})"]+[f"{short[k]}: {v}" for k,v in ex],fill=RGBColor(0xFD,0xEC,0xEE),line=RED,size=3.2,align=PP_ALIGN.LEFT)
arrow(bx+bw,3.23,rx,3.23)
arrow(bx+bw/2,3.44,bx+bw/2,3.56)
pbox(bx,3.56,bw,0.42,['Studies included in','the scoping review',f"n = {P['included']}"],fill=RGBColor(0xD9,0xE2,0xF3),line=NAVY,size=4.2)
write(tb(fx,4.04,2.47,0.26),[{'runs':[('Eligibility judged on titles, abstracts and subject terms; full-text verification ongoing. All decisions logged in the supplementary workbook.',False,True)]}],size=3.2,color=MUTED)
# inclusion criteria box (as in the earlier poster)
ib=tb(fx,4.33,2.47,0.62,fill=PALE,line=NAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,anchor=MSO_ANCHOR.MIDDLE,lw=0.5)
write(ib,[{'runs':[('Inclusion Criteria (Arksey & O’Malley, 2005 — scoping review framework)',True,False)],'align':PP_ALIGN.CENTER,'size':3.8,'space':1},
          {'runs':[('✓  Peer-reviewed empirical or qualitative studies, 2010–2025',False,False)],'align':PP_ALIGN.CENTER,'space':0.5},
          {'runs':[('✓  PMDD as the primary population',False,False)],'align':PP_ALIGN.CENTER,'space':0.5},
          {'runs':[('✓  At least one psychosocial outcome or coping mechanism as a primary aim',False,False)],'align':PP_ALIGN.CENTER,'space':0}],size=3.6,color=INK)
# --- Fig 2 domains ---
dx=5.2
write(tb(dx,fy,2.45,0.16),[[('Fig 2. Included Studies Reporting Each Outcome (of 85)',True,False)]],size=5.2,color=NAVY)
D=sorted(S['domains'].items(),key=lambda x:x[1])
lab={'Depression / depressive symptoms':'Depression','Psychological distress & perceived stress':'Psychological distress','Suicidal ideation, attempts & self-harm':'Suicidal ideation / self-harm','Interpersonal & relational functioning':'Interpersonal functioning','Academic & occupational functioning':'Academic / occupational','Lived experience, stigma & help-seeking':'Lived experience / help-seeking','Trauma, adversity & discrimination':'Trauma & adversity','Coping strategies & resilience':'Coping mechanisms','Mood & emotion regulation':'Mood & emotion regulation'}
cd=CategoryChartData(); cd.categories=[lab.get(k,k) for k,_ in D]; cd.add_series('Studies',[v for _,v in D])
def style(ch,size,maxv):
    ch.has_title=False; ch.has_legend=False; ch.font.name=FONT; ch.font.size=Pt(size); ch.font.color.rgb=INK
    pl=ch.plots[0]; pl.gap_width=40; pl.has_data_labels=True; dl=pl.data_labels; dl.font.size=Pt(size); dl.font.bold=True; dl.font.name=FONT; dl.position=XL_LABEL_POSITION.OUTSIDE_END
    ser=pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb=NAVY
    va=ch.value_axis; va.visible=False; va.has_major_gridlines=False; va.maximum_scale=maxv; va.minimum_scale=0
    ca=ch.category_axis; ca.tick_labels.font.size=Pt(size); ca.format.line.color.rgb=GRAY
    return ser
gf=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(dx),Inches(1.98),Inches(2.45),Inches(2.85),cd); ser=style(gf.chart,4.2,31)
ci=[i for i,(k,_) in enumerate(D) if k=='Coping strategies & resilience'][0]
pt=ser.points[ci]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb=RED
write(tb(dx,4.82,2.45,0.14),[{'runs':[('A study could report more than one outcome; counts should not be summed.',False,True)]}],size=3.2,color=MUTED)
# --- Fig 3 region ---
gy=5.07
write(tb(fx,gy,5.0,0.16),[[('Fig 3. Share of the 85 Included Studies, by World Bank Region and Income Group',True,False)]],size=5.2,color=NAVY)
R=sorted(S['region'].items(),key=lambda x:x[1])
rl={'Online / multinational / not reported':'Online / multinational / NR'}
cd=CategoryChartData(); cd.categories=[rl.get(k,k) for k,_ in R]; cd.add_series('Studies',[v for _,v in R])
gf=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(fx),Inches(gy+0.2),Inches(2.75),Inches(1.75),cd); ser=style(gf.chart,4.2,33)
for i,(k,_) in enumerate(R):
    if k=='Sub-Saharan Africa': p_=ser.points[i]; p_.format.fill.solid(); p_.format.fill.fore_color.rgb=RED
order=['High income','Upper-middle income','Lower-middle income','Low income','Not classifiable (online/multinational/not reported)']
names=['High','Upper-middle','Lower-middle','Low','Online / NR']
cd=CategoryChartData(); cd.categories=names; cd.add_series('Studies',[S['income'].get(k,0) for k in order])
gf=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(5.45),Inches(gy+0.2),Inches(2.2),Inches(1.75),cd); ser=style(gf.chart,4.2,65)
for i,colr in enumerate([NAVY,RED,RED,RED,GRAY]): p_=ser.points[i]; p_.format.fill.solid(); p_.format.fill.fore_color.rgb=colr
write(tb(fx,7.0,5.0,0.3),[{'runs':[('Sub-Saharan Africa: 3 studies (Ethiopia 2, Nigeria 1). Red columns = low- and middle-income countries (21 studies, 25%); only 5 (5.9%) from low or lower-middle-income countries. Income groups: World Bank FY2027.',False,True)],'align':PP_ALIGN.CENTER}],size=3.4,color=MUTED)
# ---------------- RIGHT COLUMN ----------------
RX,RW=7.773,2.117
header(RX,0.794,RW,0.36,'Results, Discussion, Conclusion, and Limitations',size=9.5)
B=lambda t:{'runs':t if isinstance(t,list) else [(t,False,False)],'bullet':True}
body(RX,1.19,RW,3.25,[
 [('Results',True,False)],
 f"Of {P['identified_total']:,} records, {P['duplicates_removed']} duplicates were removed; {P['screened']:,} titles and abstracts were screened and {P['included']} studies met the inclusion criteria (Fig. 1). Nearly half (49%) were published since 2020; most were cross-sectional (56%), and only 24% confirmed PMDD with prospective daily ratings.",
 "Depression was the most frequently examined outcome (n=28), followed by mood and emotion regulation and psychiatric comorbidity (n=18 each), psychological distress (n=16) and suicidal ideation or self-harm (n=13) (Fig. 2). PMDD was associated with suicidal ideation (OR 2.22) and attempts (OR 2.10) in a US national sample,⁵ and 72% of a global PMDD sample reported lifetime active suicidal ideation.⁶ Only 7 studies examined coping, mostly self-medication or maladaptive coping styles.",
 "Two-thirds of studies (67%) came from high-income countries; 3 came from sub-Saharan Africa (Fig. 3).",
 [('Discussion',True,False)],
 "The literature consistently links PMDD with depression, suicidality, impaired functioning and reduced quality of life, but it relies heavily on provisional diagnoses, and coping and lived experience remain understudied. No study from sub-Saharan Africa examined suicidality, quality of life or coping, so the burden there is largely undocumented.",
 [('Conclusion',True,False)],
 "PMDD research documents a substantial psychosocial burden, while coping mechanisms and low-income settings remain underrepresented. Culturally responsive screening and Africa-centred research are needed.",
 [('Limitations',True,False)],
 B([('Screening stage: ',True,False),('eligibility judged on titles and abstracts; full-text verification ongoing.',False,False)]),
 B([('Search scope: ',True,False),('English only; African Journals Online and grey literature not searched.',False,False)]),
 B([('Diagnosis: ',True,False),('many studies used provisional diagnoses or combined PMS and PMDD.',False,False)]),
 B([('Coding overlap: ',True,False),('studies can address several domains, so counts should not be summed.',False,False)])],size=4.5,align=PP_ALIGN.JUSTIFY,space=1.5)
header(RX,4.51,RW,0.36,'Public Health Implications and Future Work',size=9.5)
body(RX,4.9,RW,1.13,[
 [('Public Health Implications',True,False)],
 B([('Screen for suicide risk: ',True,False),('ask about suicidal thoughts when women present with severe premenstrual mood symptoms.',False,False)]),
 B([('Integrate care: ',True,False),('include PMDD screening in sexual and reproductive, school and university health services.',False,False)]),
 B([('Address coping needs: ',True,False),('ask how women manage symptoms and what support they need.',False,False)]),
 [('Future Work',True,False)],
 B([('Confirm PMDD prospectively ',True,False),('in studies from low- and middle-income countries.',False,False)]),
 B([('Study lived experience: ',True,False),('stigma, coping and help-seeking in African settings.',False,False)]),
 B([('Validate screening tools ',True,False),('in African languages.',False,False)])],size=4.2,space=0.9)
header(RX,6.095,RW,0.199,'References')
refs=["1. Reilly TJ et al. J Affect Disord. 2024;349:534-540.","2. American Psychiatric Association. DSM-5. 2013.","3. Arksey H, O’Malley L. Int J Soc Res Methodol. 2005;8:19-32.","4. Tricco AC et al. Ann Intern Med. 2018;169:467-473.","5. Pilver CE et al. Soc Psychiatry Psychiatr Epidemiol. 2013;48:437-446.","6. Eisenlohr-Moul T et al. BMC Psychiatry. 2022;22. doi:10.1186/s12888-022-03851-0","Scan the QR code for the full reference list."]
# move QR left, refs to right
for sh in s.shapes:
    if sh.name=='Picture 18': sh.left=Inches(RX+0.02); sh.top=Inches(6.33); sh.width=sh.height=Inches(1.0)
body(RX+1.05,6.33,RW-1.05,1.0,[{'runs':[(r,False,r.startswith('Scan'))]} for r in refs],size=3.0,space=0.6)
prs.save('PMDD_APHA_Poster_FINAL.pptx'); print('saved')
