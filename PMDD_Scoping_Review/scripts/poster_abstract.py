import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree
S=json.load(open('stats.json')); G=json.load(open('prisma_groups.json')); P=S['prisma']
NAVY=RGBColor(0x0A,0x25,0x4E); WHITE=RGBColor(255,255,255); INK=RGBColor(0x11,0x18,0x27); MUTED=RGBColor(0x4B,0x55,0x63)
RED=RGBColor(0xA3,0x1F,0x34); GRAY=RGBColor(0x9C,0xA3,0xAF); LIGHT=RGBColor(0xF3,0xF4,0xF6); PALE=RGBColor(0xEE,0xF2,0xF8); GREEN=RGBColor(0x04,0x78,0x57); DIV=RGBColor(0xD1,0xD5,0xDB)
FONT='Times New Roman'
prs=Presentation('/root/work/prev.pptx'); s=prs.slides[0]
keep={'TextBox 3','Picture 76','TextBox 304','Picture 18'}
for sh in list(s.shapes):
    if sh.name not in keep: sh._element.getparent().remove(sh._element)
# clear old title text (keep its frame)
for sh in s.shapes:
    if sh.name=='TextBox 3':
        tf=sh.text_frame
        for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
        for r in list(tf.paragraphs[0].runs): r._r.getparent().remove(r._r)
def tb(x,y,w,h,fill=None,line=None,anchor=MSO_ANCHOR.TOP,margin=0.03,shape=MSO_SHAPE.RECTANGLE,lw=0.6):
    if fill is None and line is None and shape==MSO_SHAPE.RECTANGLE: b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    else:
        b=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h)); b.shadow.inherit=False
        if fill is None: b.fill.background()
        else: b.fill.solid(); b.fill.fore_color.rgb=fill
        if line is None: b.line.fill.background()
        else: b.line.color.rgb=line; b.line.width=Pt(lw)
        if shape==MSO_SHAPE.ROUNDED_RECTANGLE: b.adjustments[0]=0.08
    tf=b.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(0.015)
    return b
def write(b,paras,size=4.4,color=INK,align=PP_ALIGN.LEFT,space=1.2,ls=None):
    tf=b.text_frame; first=True
    for it in paras:
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        if isinstance(it,dict): runs=it['runs']; bullet=it.get('bullet',False); sz=it.get('size',size); al=it.get('align',align); col=it.get('color',color); sp=it.get('space',space)
        else: runs=it; bullet=False; sz=size; al=align; col=color; sp=space
        if isinstance(runs,str): runs=[(runs,False,False)]
        p.alignment=al; p.space_after=Pt(sp)
        if ls: p.line_spacing=ls
        if bullet:
            pPr=p._p.get_or_add_pPr(); pPr.set('marL',str(int(Inches(0.08)))); pPr.set('indent',str(-int(Inches(0.07))))
            bu=etree.SubElement(pPr,qn('a:buFont')); bu.set('typeface','Arial'); bc=etree.SubElement(pPr,qn('a:buChar')); bc.set('char','•')
        for r_ in runs:
            t,bo,itl=r_[:3]; c=r_[3] if len(r_)>3 else col
            r=p.add_run(); r.text=t; f=r.font; f.name=FONT; f.size=Pt(sz); f.bold=bo; f.italic=itl; f.color.rgb=c
    return b
def header(x,y,w,h,text,size=10):
    b=tb(x,y,w,h,fill=NAVY,anchor=MSO_ANCHOR.MIDDLE); write(b,[text],size=size,color=WHITE,align=PP_ALIGN.CENTER,space=0); return b
def body(x,y,w,h,paras,size=4.4,**k):
    b=tb(x,y,w,h,fill=WHITE,margin=0.05); b.text_frame.margin_top=Inches(0.04); write(b,paras,size=size,**k); return b
# ---------------- TITLE ----------------
t=tb(1.4,0.075,8.45,0.675,anchor=MSO_ANCHOR.MIDDLE)
write(t,[{'runs':[('PSYCHOSOCIAL OUTCOMES AND COPING MECHANISMS AMONG WOMEN WITH PREMENSTRUAL DYSPHORIC DISORDER: A GLOBAL SCOPING REVIEW',True,False)],'size':10,'align':PP_ALIGN.CENTER,'space':1.5},
         {'runs':[('Ifeoluwanimi P. Shobayo, MSPH, Chelsea R. Mazonde, MPH, Marylyn O. Oduneye, MSPH, Tahirou Diallo, MSPH, Fadzai G. Nyarugwe, MSPH, Cynthia C. Ilechukwu, MSPH',False,False)],'size':7.9,'align':PP_ALIGN.CENTER,'space':0}])
# ---------------- LEFT COLUMN ----------------
LX,LW=0.107,2.217
header(LX,0.806,LW,0.199,'Abstract')
body(LX,1.03,LW,1.52,[
 [('Background: ',True,False),('Premenstrual dysphoric disorder (PMDD) is a DSM-5 depressive disorder affecting an estimated 31 million women and girls worldwide.¹ This scoping review examined psychosocial outcomes and coping mechanisms among women with PMDD and assessed the geographic distribution of the evidence.',False,False)],
 [('Methods: ',True,False),("Following Arksey and O’Malley’s framework and PRISMA-ScR guidance, we searched PubMed/MEDLINE, CINAHL, APA PsycInfo, APA PsycArticles and the Cochrane Library for peer-reviewed studies published 2010–2025. Eligible empirical and qualitative studies had PMDD as the primary population and at least one psychosocial outcome or coping mechanism as a primary aim.",False,False)],
 [('Results: ',True,False),("The search identified 1,943 records; 1,618 remained after deduplication and 44 studies met the inclusion criteria. Depression was most documented (n=39), followed by psychological distress (n=24), interpersonal functioning (n=18), suicidal ideation or self-harm (n=14) and quality of life (n=14). Only three studies examined coping mechanisms. Only three studies (6.8%) originated from low- and lower-middle-income countries, including one from Nigeria.",False,False)],
 [('Conclusions: ',True,False),('Evidence on PMDD’s psychosocial burden is growing but geographically concentrated in high-income countries, revealing a critical LMIC evidence gap.',False,False)],
 [('Keywords: ',True,False),('premenstrual dysphoric disorder; psychosocial outcomes; coping mechanisms; LMICs; scoping review',False,True)]],size=4.3,align=PP_ALIGN.JUSTIFY,space=1.2)
header(LX,2.62,LW,0.36,'Introduction, Objective, and Research Question',size=9.5)
body(LX,3.03,LW,1.36,[
 [('Introduction: ',True,False),('PMDD is a recurring, hormone-linked mood disorder in which irritability, mood swings, depressed mood and anxiety emerge in the late luteal phase and remit after menses begin. It is classified as a depressive disorder in DSM-5² and recognized in ICD-11.³ Using prospectively confirmed diagnoses, community prevalence is about 1.6%, roughly 31 million women and girls worldwide.¹ PMDD has been linked with suicidal ideation and attempts,⁴˒⁵ strained romantic relationships⁶ and reduced quality of life.⁷ Premenstrual symptoms are also common across Africa, where pooled premenstrual syndrome prevalence reaches 47%,⁸ yet it remains unclear how women with PMDD cope or where the evidence has been generated.',False,False)],
 [('Objective: ',True,False),('To map psychosocial outcomes and coping mechanisms reported among women with PMDD and identify geographic gaps in the published literature.',False,False)],
 [('Research Question: ',True,False),('What psychosocial outcomes and coping mechanisms have been reported in studies of women with PMDD published between 2010 and 2025, and how are these studies distributed across geographic settings?',False,False)]],size=4.45,align=PP_ALIGN.JUSTIFY,space=1.6)
header(LX,4.46,LW,0.199,'Methods')
rows=[('Study Design','Scoping review guided by Arksey and O’Malley’s framework⁹ with Levac et al.’s refinements,¹⁰ reported using PRISMA-ScR.¹¹'),
('Databases','PubMed/MEDLINE, CINAHL, APA PsycInfo, APA PsycArticles and Cochrane Library.'),
('Publication Years','2010–2025.'),
('Inclusion Criteria','Peer-reviewed empirical or qualitative studies with PMDD as the primary population and at least one psychosocial outcome or coping mechanism as a primary aim.'),
('Exclusion Criteria','Animal studies, laboratory research, reviews, interventions and non-peer-reviewed publications.'),
('Screening Process','1,943 records identified; 325 duplicates removed; 1,618 records screened by title and abstract; 64 assessed for eligibility; 44 studies included.'),
('Evidence Mapping','Included studies categorized by psychosocial outcome, coping mechanism, publication year and geographic setting (country and World Bank income group¹²).')]
top=4.69; H=7.36-top
gt=s.shapes.add_table(len(rows)+1,2,Inches(LX),Inches(top),Inches(LW),Inches(H)).table
gt.columns[0].width=Inches(0.62); gt.columns[1].width=Inches(LW-0.62)
gt.rows[0].height=Inches(0.14)
for i in range(1,len(rows)+1): gt.rows[i].height=Inches((H-0.14)/len(rows))
def cell(c,txt,bold=False,fill=WHITE,size=4.1):
    c.fill.solid(); c.fill.fore_color.rgb=fill; c.vertical_anchor=MSO_ANCHOR.MIDDLE
    c.margin_left=c.margin_right=Inches(0.03); c.margin_top=c.margin_bottom=Inches(0.01)
    p=c.text_frame.paragraphs[0]; p.text=''; r=p.add_run(); r.text=txt; r.font.size=Pt(size); r.font.name=FONT; r.font.bold=bold; r.font.color.rgb=INK
for j,h in enumerate(['Category','Description']): cell(gt.cell(0,j),h,True,RGBColor(0xD9,0xE2,0xF3),4.3)
for i,(a,b_) in enumerate(rows,1):
    cell(gt.cell(i,0),a,True,PALE if i%2 else WHITE); cell(gt.cell(i,1),b_,False,PALE if i%2 else WHITE)
tblPr=gt._tbl.tblPr; tblPr.set('bandRow','0')
# ---------------- CENTER PANEL ----------------
CX,CY,CW=2.536,0.794,5.179
# stat row
tiles=[('1,943','records identified','across 5 databases',NAVY),('44','studies included','in the scoping review',NAVY),
       ('3','studies addressing coping','6.8% of 44',RED),('3','low/lower-middle-income studies','6.8% of 44',RED)]
tw=(CW-0.2)/4
for k,(big,lab,sub,col) in enumerate(tiles):
    x=CX+0.1+k*tw
    b=tb(x,CY+0.07,tw,0.86,anchor=MSO_ANCHOR.MIDDLE)
    write(b,[{'runs':[(big,True,False,col)],'size':24,'align':PP_ALIGN.CENTER,'space':0},{'runs':[(lab,False,False)],'size':6,'align':PP_ALIGN.CENTER,'space':0},{'runs':[(sub,False,True,MUTED)],'size':5.5,'align':PP_ALIGN.CENTER,'space':0}])
    if k:
        ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(CY+0.15),Inches(x),Inches(CY+0.85)); ln.line.color.rgb=DIV; ln.line.width=Pt(0.5)
# --- Fig 1 PRISMA ---
fx,fy=2.64,1.78
write(tb(fx,fy,2.45,0.16),[[('Figure 1. ',True,False),('Study Selection Flow Diagram (Adapted From PRISMA-ScR¹¹)',True,False)]],size=5.2,color=NAVY)
def pbox(x,y,w,h,lines,fill=PALE,line=NAVY,size=3.6,bold_first=True,align=PP_ALIGN.CENTER):
    b=tb(x,y,w,h,fill=fill,line=line,anchor=MSO_ANCHOR.MIDDLE,shape=MSO_SHAPE.ROUNDED_RECTANGLE,margin=0.03)
    write(b,[{'runs':[(l,bold_first and k==0,False)],'align':align,'space':0} for k,l in enumerate(lines)],size=size); return b
def arrow(x1,y1,x2,y2):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); c.line.color.rgb=NAVY; c.line.width=Pt(0.9)
    ln=c.line._get_or_add_ln(); te=etree.SubElement(ln,qn('a:tailEnd')); te.set('type','triangle'); te.set('w','med'); te.set('len','med')
def stage(y,h,label):
    b=tb(fx,y,0.33,h,fill=WHITE,line=GRAY,anchor=MSO_ANCHOR.MIDDLE,shape=MSO_SHAPE.ROUNDED_RECTANGLE,margin=0.01,lw=0.4)
    write(b,[{'runs':[(label,True,True)],'align':PP_ALIGN.CENTER}],size=3.4,color=NAVY)
bx,bw=3.02,1.05; rx,rw=4.14,0.97
rowsY=[(1.98,0.4,'Identification'),(2.46,0.36,'De-duplication'),(2.9,0.4,'Screening'),(3.38,0.4,'Eligibility'),(3.86,0.4,'Included')]
for y,h,l in rowsY: stage(y,h,l)
pbox(bx,1.98,bw,0.4,['Records identified via database search','PubMed/MEDLINE · CINAHL · APA PsycInfo','APA PsycArticles · Cochrane Library','n = 1,943'],size=3.3)
arrow(bx+bw/2,2.38,bx+bw/2,2.46)
pbox(bx,2.46,bw,0.36,['Records after duplicate removal','325 duplicates removed','n = 1,618'],size=3.6)
arrow(bx+bw/2,2.82,bx+bw/2,2.9)
pbox(bx,2.9,bw,0.4,['Records screened','(title and abstract)','n = 1,618'],size=3.6)
pbox(rx,2.9,rw,0.4,['Excluded (n = 1,554)','Not PMDD-focused or no psychosocial/','coping aim; review, intervention,','animal or laboratory study'],fill=RGBColor(0xFD,0xEC,0xEE),line=RED,size=3.2)
arrow(bx+bw,3.1,rx,3.1)
arrow(bx+bw/2,3.3,bx+bw/2,3.38)
pbox(bx,3.38,bw,0.4,['Records assessed for eligibility','(title and abstract)','n = 64'],size=3.6)
pbox(rx,3.38,rw,0.4,['Excluded (n = 20)','Did not meet','inclusion criteria'],fill=RGBColor(0xFD,0xEC,0xEE),line=RED,size=3.3)
arrow(bx+bw,3.58,rx,3.58)
arrow(bx+bw/2,3.78,bx+bw/2,3.86)
pbox(bx,3.86,bw,0.4,['Studies included in','the scoping review','n = 44'],fill=RGBColor(0xD9,0xE2,0xF3),line=NAVY,size=4.2)
# inclusion criteria box (as in the earlier poster)
ib=tb(fx,4.38,2.47,0.6,fill=PALE,line=NAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,anchor=MSO_ANCHOR.MIDDLE,lw=0.5)
write(ib,[{'runs':[('Inclusion Criteria (Arksey & O’Malley, 2005 — scoping review framework)',True,False)],'align':PP_ALIGN.CENTER,'size':3.8,'space':1},
          {'runs':[('✓  Peer-reviewed empirical or qualitative studies, 2010–2025',False,False)],'align':PP_ALIGN.CENTER,'space':0.5},
          {'runs':[('✓  PMDD as the primary population',False,False)],'align':PP_ALIGN.CENTER,'space':0.5},
          {'runs':[('✓  At least one psychosocial outcome or coping mechanism as a primary aim',False,False)],'align':PP_ALIGN.CENTER,'space':0}],size=3.6,color=INK)
# --- Fig 2 domains ---
dx=5.2
write(tb(dx,fy,2.45,0.16),[[('Figure 2. ',True,False),('Psychosocial Outcomes Reported in the 44 Included Studies',True,False)]],size=5.2,color=NAVY)
D=[('Coping mechanisms',3),('Quality of life',14),('Suicidal ideation or self-harm',14),('Interpersonal functioning',18),('Psychological distress',24),('Depression',39)]
cd=CategoryChartData(); cd.categories=[k for k,_ in D]; cd.add_series('Studies',[v for _,v in D])
def style(ch,size,maxv):
    ch.has_title=False; ch.has_legend=False; ch.font.name=FONT; ch.font.size=Pt(size); ch.font.color.rgb=INK
    pl=ch.plots[0]; pl.gap_width=45; pl.has_data_labels=True; dl=pl.data_labels; dl.font.size=Pt(size); dl.font.bold=True; dl.font.name=FONT; dl.position=XL_LABEL_POSITION.OUTSIDE_END
    ser=pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb=NAVY
    va=ch.value_axis; va.visible=False; va.has_major_gridlines=False; va.maximum_scale=maxv; va.minimum_scale=0
    ca=ch.category_axis; ca.tick_labels.font.size=Pt(size); ca.format.line.color.rgb=GRAY
    return ser
gf=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(dx),Inches(1.98),Inches(2.45),Inches(2.85),cd); ser=style(gf.chart,4.6,44)
pt=ser.points[0]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb=RED
write(tb(dx,4.82,2.45,0.14),[{'runs':[('A study could report more than one outcome; counts should not be summed.',False,True)]}],size=3.2,color=MUTED)
# --- Fig 3 geography ---
gy=5.07
write(tb(fx,gy,5.0,0.16),[[('Figure 3. ',True,False),('Country Setting and World Bank Income Group of Included Studies¹²',True,False)]],size=5.2,color=NAVY)
s.shapes.add_picture('fig3_map.png',Inches(fx),Inches(gy+0.22),width=Inches(3.4))
call=tb(6.1,gy+0.35,1.5,1.2,fill=PALE,line=RED,shape=MSO_SHAPE.ROUNDED_RECTANGLE,anchor=MSO_ANCHOR.MIDDLE,lw=0.6)
write(call,[{'runs':[('6.8%',True,False,RED)],'size':20,'align':PP_ALIGN.CENTER,'space':0},{'runs':[('of included studies (3 of 44) came from low- and lower-middle-income countries',False,False)],'size':5.2,'align':PP_ALIGN.CENTER,'space':1},{'runs':[('20 high-income · 7 upper-middle-income',False,False)],'size':5.2,'align':PP_ALIGN.CENTER,'space':1},{'runs':[('Sub-Saharan Africa: 1 study (Nigeria)',True,False,RED)],'size':5.2,'align':PP_ALIGN.CENTER,'space':0}])
write(tb(fx,6.82,5.0,0.22),[{'runs':[('Map shows the 30 studies with an identifiable country setting (World Bank income groups). The other 14 used online or multinational samples or had no single mappable setting in the available records. *Includes Hong Kong (1).',False,True)],'align':PP_ALIGN.CENTER}],size=4.2,color=MUTED)
# ---------------- RIGHT COLUMN ----------------
RX,RW=7.773,2.117
header(RX,0.794,RW,0.36,'Results, Discussion, Conclusion, and Limitations',size=9.5)
B=lambda t:{'runs':t if isinstance(t,list) else [(t,False,False)],'bullet':True}
body(RX,1.19,RW,3.25,[
 [('Results',True,False)],
 "The database search identified 1,943 records. After 325 duplicates were removed, 1,618 records were screened by title and abstract; 64 were assessed for eligibility, 20 were excluded, and 44 studies were included (Figure 1). Most (61%) were published in 2022–2025 (Figure 4).",
 "Depression was the most documented outcome (n=39); one nationwide cohort found a 2.6-fold higher risk of later depression after a PMDD diagnosis.¹³ Psychological distress (n=24), interpersonal functioning (n=18), suicidal ideation or self-harm (n=14) and quality of life (n=14) followed. Only three studies examined coping mechanisms, such as peer support shared in online PMDD communities¹⁴ (Figure 2).",
 "Twenty studies came from high-income countries and seven from upper-middle-income countries (Türkiye, Iran); only three (6.8%) came from low- and lower-middle-income countries—Bangladesh,¹⁵ Lebanon¹⁶ and Nigeria,¹⁷ the only study from sub-Saharan Africa (Figure 3).",
 [('Discussion',True,False)],
 "PMDD research consistently documents depression, distress, relationship difficulties and suicidality.⁴⁻⁷ Patients describe misdiagnosis and uneven provider knowledge,¹⁸˒¹⁹ yet coping remains largely unexamined, and the scarcity of studies from low-income settings and Africa leaves the burden there invisible to health systems.",
 [('Conclusion',True,False)],
 "Evidence on PMDD’s psychosocial burden is growing but geographically concentrated in high-income countries, calling for culturally responsive screening, integrated mental and reproductive health services, and Africa-centred research investment.",
 [('Limitations',True,False)],
 B([('Search scope: ',True,False),('studies outside the five databases or the 2010–2025 period may have been missed.',False,False)]),
 B([('Study differences: ',True,False),('variation in design and outcome definitions limits direct comparison.',False,False)]),
 B([('Coding overlap: ',True,False),('a study could address more than one outcome, so counts should not be summed.',False,False)]),
 B([('Geographic coverage: ',True,False),('few studies from low- and lower-middle-income countries limit conclusions about those settings.',False,False)])],size=4.1,align=PP_ALIGN.JUSTIFY,space=1.0)
header(RX,4.51,RW,0.36,'Public Health Implications and Future Work',size=9.5)
body(RX,4.9,RW,1.13,[
 [('Public Health Implications',True,False)],
 B([('Screen for suicide risk: ',True,False),('ask about suicidal thoughts when women present with severe premenstrual mood symptoms.⁴˒⁵',False,False)]),
 B([('Integrate care: ',True,False),('embed PMDD screening in sexual and reproductive, school and university health services.',False,False)]),
 B([('Train providers: ',True,False),('improve recognition to reduce misdiagnosis and delayed care.¹⁸',False,False)]),
 B([('Address coping needs: ',True,False),('ask how women manage symptoms and connect them with peer and professional support.¹⁴',False,False)]),
 [('Future Work',True,False)],
 B([('Confirm PMDD prospectively ',True,False),('using validated daily ratings.¹˒²⁰',False,False)]),
 B([('Prioritize low-income and African settings, ',True,False),('including qualitative studies of stigma and coping.',False,False)]),
 B([('Validate screening tools ',True,False),('in African languages and test coping-focused interventions.',False,False)])],size=4.05,space=0.7)
header(RX,6.095,RW,0.199,'References')
refs=["1. Reilly TJ, et al. J Affect Disord. 2024;349:534-540.","2. American Psychiatric Association. DSM-5. 2013.","3. World Health Organization. ICD-11. 2022.","4. Pilver CE, et al. Soc Psychiatry Psychiatr Epidemiol. 2013;48:437-446.","5. Eisenlohr-Moul T, et al. BMC Psychiatry. 2022;22:199.","6. Westermark V, et al. J Affect Disord. 2024;364:132-138.","7. Wang Q, et al. JAMA Netw Open. 2025;8:e2533823.","8. Andualem F, et al. Front Psychiatry. 2024;15:1338304.","9. Arksey H, O’Malley L. Int J Soc Res Methodol. 2005;8:19-32.","10. Levac D, et al. Implement Sci. 2010;5:69.","11. Tricco AC, et al. Ann Intern Med. 2018;169:467-473.","12. World Bank. Country income classifications, FY2027. 2026.","13. Li DJ, et al. Asian J Psychiatr. 2023;79:103355.","14. Winslow A, et al. Womens Reprod Health. 2023;10:420-435.","15. Roy N, et al. PLoS One. 2025;20:e0321097.","16. Younes Y, et al. BMC Psychiatry. 2021;21:548.","17. Adegoke AA, et al. Gend Behav. 2014;12:6087-6094.","18. Hantsoo L, et al. J Womens Health. 2022;31:100-109.","19. Habib S, et al. Womens Reprod Health. 2025;12:310-329.","20. Naik SS, et al. Front Glob Womens Health. 2023;4:1181583."]
for sh in s.shapes:
    if sh.name=='Picture 18': sh.left=Inches(8.272); sh.top=Inches(6.327); sh.width=sh.height=Inches(1.129)
# key takeaway banner (center)
kb=tb(CX+0.12,7.06,CW-0.24,0.26,fill=NAVY,anchor=MSO_ANCHOR.MIDDLE,shape=MSO_SHAPE.ROUNDED_RECTANGLE,margin=0.08)
write(kb,[{'runs':[('KEY TAKEAWAY  ',True,False,RGBColor(0xF2,0xC1,0x4E)),('PMDD’s psychosocial burden is well documented in high-income countries, yet nearly invisible in low-income settings and sub-Saharan Africa.',False,False,WHITE)],'align':PP_ALIGN.CENTER}],size=5.6)
write(tb(RX+0.06,3.62,RW-0.12,0.13),[[('Figure 4. ',True,False),('Included Studies by Publication Period',True,False)]],size=4.6,color=NAVY)
cd=CategoryChartData(); cd.categories=['2010–13','2014–17','2018–21','2022–25']; cd.add_series('Studies',[6,6,5,27])
gf=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(RX+0.1),Inches(3.74),Inches(RW-0.2),Inches(0.58),cd); ser=style(gf.chart,4.2,32)
p_=ser.points[3]; p_.format.fill.solid(); p_.format.fill.fore_color.rgb=RED
write(tb(RX+0.06,4.32,RW-0.12,0.1),[{'runs':[('27 of 44 studies (61%) were published in 2022–2025.',False,True)],'align':PP_ALIGN.CENTER}],size=3.5,color=MUTED)
prs.save('PMDD_APHA_Poster_FINAL.pptx'); print('saved')
