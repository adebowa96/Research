"""Title/abstract re-screening of the 516 exported records (324 unique after deduplication).

Stage 1 applies objective rules (publication year, article type). Stage 2 records
title/abstract decisions made by reading each remaining record against the review's criteria.
These are AI-ASSISTED PROPOSALS: every decision must be checked by the human screeners
(columns "Screener 1 decision" / "Screener 2 decision" in the workbook) before it is final.

Eligibility criteria applied
  Include: primary quantitative, qualitative, or mixed-methods study; published Jan 2019 – Mar
  2026; population includes individuals with MRKH (results reported for them); reports at least
  one psychological / mental health outcome (depression, anxiety, distress, self-esteem, body or
  genital image, quality of life, psychosexual wellbeing) or coping / psychosocial experiences.
  Exclude:
    E1 published before 2019
    E2 not primary research (review, commentary, letter, erratum, book chapter, educational piece)
    E3 case report, small surgical case series, or surgical video
    E4 wrong population (not MRKH, or MRKH results not reported separately) or unrelated topic
    E5 anatomical, surgical, functional, or fertility outcomes only (no psychological outcome)
  Uncertain: cannot be decided from title/abstract; needs full text.
"""
import csv
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
REASONS = {
    "E1": "E1 Published before 2019",
    "E2": "E2 Not primary research (review, commentary, letter, erratum, chapter)",
    "E3": "E3 Case report / small surgical case series / surgical video",
    "E4": "E4 Wrong population or unrelated topic",
    "E5": "E5 Anatomical, surgical, functional or fertility outcomes only",
}

# Stage 2: decisions from reading titles/abstracts (record_id -> (decision, note))
INCLUDE = {
    "R002": "Case series with standardized self-esteem, depression, anxiety measures",
    "R003": "Case-control: sexual esteem, genital self-image, psychological functioning",
    "R004": "Prevalence of anxiety and depression (Malaysia)",
    "R007": "Quality of life and sexuality after surgery (French-language article)",
    "R010": "Qualitative: illness experience and unmet needs (online forum)",
    "R011": "Qualitative phenomenological study of disease-related experiences",
    "R012": "Cross-sectional: quality of life and sexual function",
    "R013": "IPA: identity, perceptions and isolation",
    "R014": "Qualitative life-course study of living with MRKH",
    "R017": "Qualitative: impact on sexual wellbeing",
    "R020": "Psychological evaluation of MRKH uterus transplant candidates",
    "R021": "Long-term mental condition and quality of life after neovagina",
    "R032": "Qualitative: coping strategies",
    "R034": "Descriptive study: sexual function and quality of life",
    "R038": "Prospective psychological intervention study",
    "R040": "Interview study: vaginal lengthening and sexual wellbeing",
    "R052": "Retrospective cohort: psychiatric comorbidities in Mullerian aplasia",
    "R057": "Qualitative: perceptions of surrogacy (psychosocial experience)",
    "R060": "Body image before and after vaginoplasty",
    "R063": "Qualitative: sexual well-being, genital self-image, coping",
    "R065": "Pre-post: distress, depression, QoL; support intervention (German-language)",
    "R073": "Pilot survey: healthcare experiences",
    "R084": "Cross-sectional: anxiety symptoms",
    "R086": "Case-control: sexual and psychosocial outcomes (Scopus labels it a note)",
    "R091": "Sexual self-esteem and psychological correlates",
    "R094": "Qualitative: sexuality of adolescents and young women",
    "R104": "Depressive symptoms in 141 patients",
    "R105": "Illness representations, self-concept and psychological adjustment",
    "R112": "Mixed methods: experiences of vaginal lengthening (n = 616)",
    "R131": "RCT: psychosexual education, sexual distress, genital self-image",
    "R149": "Qualitative with quantitative component: effect of diagnosis (Malaysia)",
    "R164": "Qualitative/quantitative: sexual identity (French-language article)",
    "R227": "Low self-esteem in MRKH (no abstract in export; confirm at full text)",
    "R267": "Qualitative: diagnostic odyssey (Denmark)",
}
UNCERTAIN = {
    "R001": "Uterine factor infertility sample (MRKH + hysterectomy): check MRKH results reported",
    "R005": "Uterus transplant candidates with uterine factor infertility: check MRKH subgroup",
    "R025": "DSD sample with 5 conditions: check whether MRKH participants are reported",
    "R028": "Vaginal aplasia/hypoplasia sample: check MRKH proportion and psychosocial outcomes",
    "R093": "Survey of attitudes to uterus transplantation: psychological outcome unclear",
    "R173": "Psychodynamic analysis (French); partly case-based: check design",
    "R175": "Sexual experience before treatment (MRKH + CAIS): check psychosexual outcomes",
    "R223": "Experience of medical encounter in Africa: no abstract in export",
    "R234": "Uterus transplant recipients' self-image: no abstract; check population",
    "R242": "Dilation vs surgery: no abstract; check for QoL/psychological outcomes",
    "R244": "Psychological assessment of uterus transplant candidates: no abstract",
    "R255": "Disclosure and stigma in adults with DSD: check MRKH subgroup",
}
MANUAL_EXCLUDE = {
    "R009": "E5", "R015": "E5", "R024": "E3", "R026": "E4", "R029": "E3", "R033": "E5",
    "R035": "E5", "R037": "E5", "R039": "E3", "R044": "E5", "R049": "E3", "R056": "E4",
    "R074": "E4", "R075": "E5", "R078": "E2", "R081": "E5", "R082": "E3", "R085": "E5",
    "R089": "E5", "R090": "E4", "R095": "E5", "R101": "E2", "R109": "E4", "R113": "E5",
    "R115": "E5", "R117": "E5", "R120": "E5", "R122": "E5", "R124": "E4", "R125": "E5",
    "R128": "E4", "R130": "E4", "R135": "E5", "R136": "E3", "R143": "E5", "R153": "E4",
    "R156": "E3", "R161": "E4", "R207": "E5", "R208": "E5", "R210": "E5", "R215": "E3",
    "R220": "E2", "R225": "E4", "R233": "E4", "R236": "E5", "R237": "E5", "R238": "E5",
    "R243": "E5", "R245": "E5", "R248": "E3", "R249": "E5", "R250": "E5", "R257": "E2",
    "R264": "E4", "R265": "E4", "R266": "E3", "R268": "E4", "R269": "E2", "R270": "E2",
    "R273": "E3", "R274": "E3", "R275": "E3", "R277": "E3", "R278": "E5", "R279": "E3",
    "R280": "E2", "R284": "E3", "R285": "E4", "R286": "E2", "R287": "E5", "R288": "E4",
    "R291": "E3", "R293": "E4", "R295": "E5", "R296": "E4", "R297": "E4", "R300": "E5",
    "R301": "E5", "R302": "E5", "R303": "E4", "R306": "E5", "R308": "E4", "R310": "E5",
    "R311": "E4", "R312": "E5", "R313": "E3", "R314": "E4", "R315": "E4", "R316": "E4",
    "R317": "E4", "R318": "E3", "R320": "E4", "R323": "E3", "R324": "E3",
}


def stage1(x):
    y = int(x["year"]) if x["year"].isdigit() else 0
    t, d = x["title"].lower(), x["doctype"].lower()
    if y and y < 2019:
        return "E1"
    if re.search(r"\breview\b|meta-analys|editorial|letter|erratum|comment|\bnote\b|short survey|book chapter", d):
        if not (re.search(r"journal article$|^article$", d) and not re.search(r"review|meta", d)):
            return "E2"
    if re.search(r"systematic review|meta-analysis|narrative review|scoping review|literature review|"
                 r"\ban update\b|comprehensive update|review and update|overview", t):
        return "E2"
    if re.search(r"case report|a case of|a rare case|case series|\bcase\b.*report|first recorded case|"
                 r"video|tutorial", t) or "case report" in d:
        return "E3"
    return ""


unique = list(csv.DictReader(open(HERE / "records_unique.csv", encoding="utf-8")))
for x in unique:
    rid = x["record_id"]
    if rid in INCLUDE:
        x["stage"], x["decision"], x["reason"], x["note"] = "2 (reviewer read)", "Include", "", INCLUDE[rid]
    elif rid in UNCERTAIN:
        x["stage"], x["decision"], x["reason"], x["note"] = "2 (reviewer read)", "Uncertain", "", UNCERTAIN[rid]
    elif rid in MANUAL_EXCLUDE:
        x["stage"], x["decision"], x["reason"], x["note"] = "2 (reviewer read)", "Exclude", REASONS[MANUAL_EXCLUDE[rid]], ""
    else:
        code = stage1(x)
        assert code, f"{rid} has no decision: {x['title']}"
        x["stage"], x["decision"], x["reason"], x["note"] = "1 (rule)", "Exclude", REASONS[code], ""
    lang = x["language"].lower()
    x["non_english"] = "yes" if lang and not lang.startswith("eng") else ""

counts = Counter(x["decision"] for x in unique)
reasons = Counter(x["reason"] for x in unique if x["decision"] == "Exclude")
all_rows = list(csv.DictReader(open(HERE / "records_all.csv", encoding="utf-8")))
summary = [
    ("Records identified (4 export files)", len(all_rows)),
    ("Duplicates removed (DOI or title match)", len(all_rows) - len(unique)),
    ("Unique records screened (title/abstract)", len(unique)),
    ("Excluded at title/abstract", counts["Exclude"]),
] + [(f"   {r}", n) for r, n in sorted(reasons.items())] + [
    ("Uncertain: need full text", counts["Uncertain"]),
    ("Proposed include (full text to confirm)", counts["Include"]),
]
for label, n in summary:
    print(f"{label:<75}{n}")

# ------------------------------------------------------------------ workbook
wb = Workbook()
ws = wb.active
ws.title = "Screening"
cols = ["record_id", "decision", "reason", "note", "stage", "Screener 1 decision",
        "Screener 2 decision", "Final decision", "year", "title", "authors", "journal", "doi",
        "pmid", "sources", "doctype", "language", "non_english", "abstract"]
ws.append(cols)
fills = {"Include": "D9EAD3", "Uncertain": "FFF2CC", "Exclude": "F4CCCC"}
order = {"Include": 0, "Uncertain": 1, "Exclude": 2}
for x in sorted(unique, key=lambda x: (order[x["decision"]], x["record_id"])):
    ws.append([x.get(c, "") for c in cols])
    ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=fills[x["decision"]])
for c in ws[1]:
    c.font = Font(bold=True)
widths = dict(record_id=9, decision=11, reason=34, note=44, stage=14, title=70, abstract=60)
for i, c in enumerate(cols, 1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 16 if "Screener" in c or c == "Final decision" else 14)
ws.freeze_panes = "C2"
ws.auto_filter.ref = ws.dimensions
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=False, vertical="top")

s2 = wb.create_sheet("PRISMA counts")
s2.append(["Stage", "n"])
for label, n in summary:
    s2.append([label, n])
s2.column_dimensions["A"].width = 70
s2.append([])
s2.append(["These counts are provisional until both screeners confirm every decision and "
           "full-text review of Include + Uncertain records is complete."])

s3 = wb.create_sheet("Criteria & instructions")
for line in __doc__.strip().splitlines():
    s3.append([line])
s3.append([])
for line in [
    "HOW TO USE THIS SHEET",
    "1. Two people (e.g., you and Dr. Okojie or Dr. Anderson) each fill in 'Screener 1 decision' "
    "and 'Screener 2 decision' independently for every record, without looking at each other's.",
    "2. Where screeners disagree, discuss and record the agreed 'Final decision'.",
    "3. Obtain full texts for every record whose Final decision is Include or Uncertain and apply "
    "the same criteria to the full text; record reasons for any full-text exclusion.",
    "4. Send the completed workbook back so the PRISMA flow, results, and appendix can be rebuilt.",
    "Column 'decision' is an AI-assisted proposal from title/abstract only and is not final.",
]:
    s3.append([line])
s3.column_dimensions["A"].width = 120

dups = wb.create_sheet("All 516 records")
dcols = ["record_id", "duplicate", "source", "database", "year", "title", "doi", "pmid"]
dups.append(dcols)
for r in all_rows:
    dups.append([r.get(c, "") for c in dcols])
dups.column_dimensions["F"].width = 90

out = HERE / "MRKH_rescreening_workbook.xlsx"
wb.save(out)
print("wrote", out)
